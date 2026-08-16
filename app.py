import datetime as dt
import json
import os
import random
import secrets
import sys
import time
from functools import wraps

sys.stdout.reconfigure(encoding="utf-8")

from flask import Flask, abort, flash, jsonify, redirect, render_template, request, url_for
from flask_login import LoginManager, current_user, login_required, login_user, logout_user
from flask_wtf.csrf import CSRFProtect

from database import db
from models import (
    User, Course, Schedule, Enrollment, Attendance,
    TeacherPayroll, TeachingRecord, TuitionPayment,
    Exam, ExamQuestion, ExamAnswer, ExamSubmission, SubmissionAnswer,
    QuestionBank, BankQuestion, BankQuestionOption, ImportSession, StudentPerformance
)
from services.import_service import ImportService
from adaptive_engine import AdaptiveEngine
from forms import RegistrationForm, LoginForm, UpdateProfileForm, ChangePasswordForm
from helpers import (
    WEEKDAYS, STATUS_LABELS, TUITION_LABELS, ROLE_LABELS, SUBJECT_COLORS,
    save_picture, role_required, manager_or_admin_required, teacher_manager_admin_required,
    parse_date, parse_time, money, duration_hours, resolve_student_id, month_key,
    week_bounds, month_bounds, get_accessible_student_ids, schedules_for_user,
    schedule_occurrences, students_for_teacher, matching_schedule, rebuild_payroll_month,
    sync_teaching_record, teaching_summary, teacher_dashboard_data, student_dashboard_data,
    current_month_summary, build_timetable_grid, get_course_color_map
)


from whitenoise import WhiteNoise

basedir = os.path.abspath(os.path.dirname(__file__))
app = Flask(__name__)
app.wsgi_app = WhiteNoise(app.wsgi_app, root=os.path.join(basedir, 'static'), prefix='static/')
app.config["SECRET_KEY"] = "khoa-bi-mat-sieu-cap-vipro-123456"

# Check if running on Vercel
IS_VERCEL = os.environ.get("VERCEL") == "1" or os.environ.get("VERCEL") is not None

# Configure database connection
DATABASE_URL = os.environ.get("DATABASE_URL") or os.environ.get("POSTGRES_URL")

if DATABASE_URL:
    if DATABASE_URL.startswith("postgres://"):
        DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)
    app.config["SQLALCHEMY_DATABASE_URI"] = DATABASE_URL
    if DATABASE_URL.startswith("postgresql"):
        app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
            "pool_size": 10,
            "max_overflow": 20,
            "pool_recycle": 300,
            "pool_pre_ping": True,
        }
else:
    if IS_VERCEL:
        db_path = "/tmp/tutoring_center.db"
        original_db_path = os.path.join(basedir, "tutoring_center.db")
        if os.path.exists(original_db_path) and not os.path.exists(db_path):
            import shutil
            shutil.copy(original_db_path, db_path)
        app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + db_path
    else:
        app.config["SQLALCHEMY_DATABASE_URI"] = "sqlite:///" + os.path.join(basedir, "tutoring_center.db")

if IS_VERCEL:
    app.config["UPLOAD_FOLDER_AVATARS"] = "/tmp/static/avatars"
    app.config["UPLOAD_FOLDER_COURSES"] = "/tmp/static/course_images"
else:
    app.config["UPLOAD_FOLDER_AVATARS"] = os.path.join(basedir, "static", "avatars")
    app.config["UPLOAD_FOLDER_COURSES"] = os.path.join(basedir, "static", "course_images")

app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Ensure upload directories exist
os.makedirs(app.config["UPLOAD_FOLDER_AVATARS"], exist_ok=True)
os.makedirs(app.config["UPLOAD_FOLDER_COURSES"], exist_ok=True)

from database import db
db.init_app(app)
csrf = CSRFProtect(app)
login_manager = LoginManager(app)
login_manager.login_view = "login"
login_manager.login_message = "Vui lòng đăng nhập để sử dụng tính năng này."
login_manager.login_message_category = "info"


def serve_static_upload(folder_key, filename, fallback_subfolder):
    from flask import send_from_directory
    folder = app.config[folder_key]
    if IS_VERCEL and os.path.exists(os.path.join(folder, filename)):
        return send_from_directory(folder, filename)
    return send_from_directory(os.path.join(basedir, "static", fallback_subfolder), filename)


@app.route('/static/avatars/<path:filename>')
def serve_avatar(filename):
    return serve_static_upload("UPLOAD_FOLDER_AVATARS", filename, "avatars")


@app.route('/static/course_images/<path:filename>')
def serve_course_image(filename):
    return serve_static_upload("UPLOAD_FOLDER_COURSES", filename, "course_images")


@login_manager.user_loader
def load_user(user_id):
    return db.session.get(User, int(user_id))


@app.context_processor
def inject_globals():
    import datetime as dt
    today = dt.date.today()
    return {
        "WEEKDAYS": dict(WEEKDAYS),
        "STATUS_LABELS": STATUS_LABELS,
        "TUITION_LABELS": TUITION_LABELS,
        "ROLE_LABELS": ROLE_LABELS,
        "current_month_summary": current_month_summary,
        "money": lambda value: f"{int(value or 0):,}".replace(",", "."),
        "today_date": today.strftime('%d/%m/%Y'),
        "now": dt.datetime.now(),
    }


# ──────────────────────── AUTH ROUTES ────────────────────────

@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("index"))
    form = RegistrationForm()
    if form.validate_on_submit():
        if form.role.data in ("admin", "manager"):
            flash("Không thể đăng ký vai trò này.", "danger")
            return redirect(url_for("register"))
        user = User(
            username=form.username.data,
            fullname=form.fullname.data,
            user_code=form.user_code.data,
            role=form.role.data,
            phone=form.phone.data,
            parent_phone=form.parent_phone.data,
            position=form.position.data,
            is_active=True,
        )
        user.set_password(form.password.data)
        db.session.add(user)
        db.session.commit()
        flash("Đăng ký tài khoản thành công.", "success")
        return redirect(url_for("login"))
    return render_template("register.html", form=form)


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("index"))
    form = LoginForm()
    if form.validate_on_submit():
        user = User.query.filter_by(username=form.username.data).first()
        if user and user.check_password(form.password.data) and user.is_active:
            login_user(user)
            return redirect(request.args.get("next") or url_for("index"))
        flash("Sai tên đăng nhập, mật khẩu hoặc tài khoản đã bị khóa.", "danger")
    return render_template("login.html", form=form)


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/profile", methods=["GET", "POST"])
@login_required
def profile():
    profile_form = UpdateProfileForm()
    password_form = ChangePasswordForm()

    if profile_form.submit_profile.data and profile_form.validate():
        if profile_form.avatar.data:
            current_user.avatar = save_picture(profile_form.avatar.data, app.config["UPLOAD_FOLDER_AVATARS"])
        current_user.fullname = profile_form.fullname.data
        current_user.username = profile_form.username.data
        current_user.phone = profile_form.phone.data
        current_user.parent_phone = profile_form.parent_phone.data if current_user.role == "student" else None
        current_user.position = profile_form.position.data
        current_user.birth_date = profile_form.birth_date.data
        db.session.commit()
        flash("Cập nhật hồ sơ thành công.", "success")
        return redirect(url_for("profile"))

    if password_form.submit_password.data and password_form.validate():
        if not current_user.check_password(password_form.old_password.data):
            flash("Mật khẩu cũ không đúng.", "danger")
        else:
            current_user.set_password(password_form.new_password.data)
            db.session.commit()
            flash("Đổi mật khẩu thành công.", "success")
            return redirect(url_for("profile"))

    if request.method == "GET":
        profile_form.fullname.data = current_user.fullname
        profile_form.username.data = current_user.username
        profile_form.phone.data = current_user.phone
        profile_form.parent_phone.data = current_user.parent_phone
        profile_form.position.data = current_user.position
        profile_form.birth_date.data = current_user.birth_date

    image_file = url_for("static", filename="avatars/" + (current_user.avatar or "default.jpg"))
    return render_template("profile.html", profile_form=profile_form, password_form=password_form, image_file=image_file)


# ──────────────────────── DASHBOARD ────────────────────────

@app.route("/")
@login_required
def index():
    today = dt.date.today()

    # Default variables to prevent any jinja2.exceptions.UndefinedError
    defaults = {
        "dashboard_scope": "center",
        "student": None,
        "teacher_data": None,
        "student_data": None,
        "unpaid_tuition": 0,
        "timetable_grid": None,
        "anchor": today,
        "recent_attendance": [],
        "payroll_history": [],
        "exams_data": [],
        "activities": [],
        "unpaid_payments": [],
        "recent_payments": [],
        "recent_teaching_records": [],
        "schedules": [],
        "total_students": 0,
        "total_teachers": 0,
        "monthly_revenue": 0,
        "unpaid_tuition_count": 0,
        "total_teaching_hours": 0,
        "total_teacher_payroll": 0
    }

    if current_user.role == "teacher":
        teacher_data = teacher_dashboard_data(current_user)
        user_schedules = schedules_for_user(current_user)
        start_week, end_week = week_bounds(today)
        timetable_grid = build_timetable_grid(user_schedules, start_week, end_week)
        
        # Recent attendance
        course_ids = [c.id for c in Course.query.filter_by(teacher_id=current_user.id).all()]
        recent_attendance = (
            Attendance.query.filter(Attendance.course_id.in_(course_ids))
            .order_by(Attendance.date.desc(), Attendance.created_at.desc())
            .limit(10).all()
        ) if course_ids else []

        payroll_history = (
            TeacherPayroll.query.filter_by(teacher_id=current_user.id)
            .order_by(TeacherPayroll.month.desc())
            .limit(5).all()
        )

        teacher_exams = Exam.query.filter_by(created_by=current_user.id).order_by(Exam.created_at.desc()).all()
        exams_data = []
        for exam in teacher_exams:
            pending_grading = ExamSubmission.query.filter_by(exam_id=exam.id, is_graded=False).count()
            total_subs = ExamSubmission.query.filter_by(exam_id=exam.id).count()
            exams_data.append({
                "exam": exam,
                "pending_grading": pending_grading,
                "total_submissions": total_subs
            })

        # Activity feed
        activities = []
        exam_ids = [e.id for e in teacher_exams]
        if exam_ids:
            subs = ExamSubmission.query.filter(ExamSubmission.exam_id.in_(exam_ids)).order_by(ExamSubmission.submitted_at.desc()).limit(10).all()
            for s in subs:
                activities.append({
                    "type": "exam_submission",
                    "title": f"Nộp bài kiểm tra: {s.exam.title}",
                    "desc": f"Học sinh {s.student.fullname} đã nộp bài. Trạng thái: {'Đã chấm' if s.is_graded else 'Chờ chấm'}",
                    "time": s.submitted_at,
                    "icon": "file-check"
                })
        records = TeachingRecord.query.filter_by(teacher_id=current_user.id).order_by(TeachingRecord.date.desc()).limit(10).all()
        for r in records:
            activities.append({
                "type": "teaching_record",
                "title": f"Báo giảng ngày {r.date.strftime('%d/%m/%Y')}",
                "desc": f"Khóa {r.course.course_name} ({r.hours_taught}h) - {'ĐÃ XÁC NHẬN' if r.confirmed_by_teacher else 'CHỜ XÁC NHẬN'}",
                "time": dt.datetime.combine(r.date, dt.time.min),
                "icon": "clipboard-check"
            })
        activities.sort(key=lambda x: x["time"], reverse=True)
        activities = activities[:8]

        context = {
            "dashboard_scope": "teacher",
            "teacher_data": teacher_data,
            "timetable_grid": timetable_grid,
            "recent_attendance": recent_attendance,
            "payroll_history": payroll_history,
            "exams_data": exams_data,
            "activities": activities
        }
        return render_template("index.html", **{**defaults, **context})

    if current_user.role in ("student", "parent"):
        student_id = current_user.id if current_user.role == "student" else current_user.linked_student_id
        student = db.session.get(User, student_id) if student_id else None
        
        if student:
            student_data = student_dashboard_data(student)
            
            unpaid_tuition = (
                db.session.query(db.func.sum(TuitionPayment.amount))
                .filter(
                    TuitionPayment.student_id == student.id,
                    TuitionPayment.status.in_(["unpaid", "partial"])
                ).scalar() or 0
            )

            user_schedules = schedules_for_user(student)
            start_week, end_week = week_bounds(today)
            timetable_grid = build_timetable_grid(user_schedules, start_week, end_week)

            unpaid_payments = TuitionPayment.query.filter(
                TuitionPayment.student_id == student.id,
                TuitionPayment.status.in_(["unpaid", "partial"])
            ).order_by(TuitionPayment.payment_date.asc()).all()

            course_ids = [e.course_id for e in Enrollment.query.filter_by(student_id=student.id, status="active").all()]
            student_exams = Exam.query.filter(Exam.course_id.in_(course_ids), Exam.is_active == True).order_by(Exam.created_at.desc()).all() if course_ids else []
            exams_data = []
            for exam in student_exams:
                submission = ExamSubmission.query.filter_by(exam_id=exam.id, student_id=student.id).first()
                status = "not_started"
                score_str = ""
                if submission:
                    status = "graded" if submission.is_graded else "submitted"
                    score_str = f"{submission.total_score}/{submission.max_score}"
                exams_data.append({
                    "exam": exam,
                    "status": status,
                    "score_str": score_str,
                    "submission_id": submission.id if submission else None
                })

            activities = []
            att_records = Attendance.query.filter_by(student_id=student.id).order_by(Attendance.date.desc()).limit(10).all()
            for att in att_records:
                status_vi = "Có mặt" if att.status == "present" else "Vắng mặt" if att.status == "absent" else "Đi trễ"
                activities.append({
                    "type": "attendance",
                    "title": f"Điểm danh: {status_vi}",
                    "desc": f"Khóa {att.course.course_name} ngày {att.date.strftime('%d/%m/%Y')}",
                    "time": dt.datetime.combine(att.date, dt.time.min),
                    "icon": "user-check"
                })
            payments = TuitionPayment.query.filter_by(student_id=student.id).order_by(TuitionPayment.payment_date.desc()).limit(10).all()
            for p in payments:
                activities.append({
                    "type": "payment",
                    "title": f"Đóng học phí: {p.amount:,} đ",
                    "desc": f"Khóa {p.course.course_name} ({p.status})",
                    "time": dt.datetime.combine(p.payment_date, dt.time.min),
                    "icon": "coins"
                })
            subs = ExamSubmission.query.filter_by(student_id=student.id).order_by(ExamSubmission.submitted_at.desc()).limit(10).all()
            for s in subs:
                desc = f"Điểm số: {s.total_score}/{s.max_score}" if s.is_graded else "Đang chờ chấm"
                activities.append({
                    "type": "exam",
                    "title": f"Nộp bài kiểm tra: {s.exam.title}",
                    "desc": desc,
                    "time": s.submitted_at,
                    "icon": "file-question"
                })
            activities.sort(key=lambda x: x["time"], reverse=True)
            activities = activities[:8]

            context = {
                "dashboard_scope": "student",
                "student_data": student_data,
                "student": student,
                "unpaid_tuition": unpaid_tuition,
                "timetable_grid": timetable_grid,
                "unpaid_payments": unpaid_payments,
                "exams_data": exams_data,
                "activities": activities
            }
        else:
            context = {
                "dashboard_scope": "student",
                "student_data": None,
                "student": None
            }
        return render_template("index.html", **{**defaults, **context})

    # Admin / Manager dashboard
    month_start = today.replace(day=1)
    total_students = User.query.filter_by(role="student").count()
    total_teachers = User.query.filter_by(role="teacher").count()
    total_courses = Course.query.count()
    total_classes_today = Schedule.query.filter_by(weekday=today.weekday()).count()
    monthly_revenue = (
        db.session.query(db.func.coalesce(db.func.sum(TuitionPayment.amount), 0))
        .filter(TuitionPayment.payment_date >= month_start, TuitionPayment.status.in_(["paid", "partial"]))
        .scalar()
    )
    unpaid_tuition_count = TuitionPayment.query.filter(TuitionPayment.status.in_(["unpaid", "partial"])).count()
    teachers_list = User.query.filter_by(role="teacher").all()
    payroll_rows = [current_month_summary(teacher.id) for teacher in teachers_list]
    total_teaching_hours = sum(row["total_hours"] for row in payroll_rows)
    total_teacher_payroll = sum(row["salary"] for row in payroll_rows)

    all_schedules = Schedule.query.order_by(Schedule.weekday, Schedule.start_time).all()
    start_week, end_week = week_bounds(today)
    timetable_grid = build_timetable_grid(all_schedules, start_week, end_week)

    recent_attendance = Attendance.query.order_by(Attendance.date.desc(), Attendance.created_at.desc()).limit(15).all()

    recent_payments = TuitionPayment.query.order_by(TuitionPayment.payment_date.desc()).limit(6).all()
    recent_teaching_records = TeachingRecord.query.order_by(TeachingRecord.date.desc()).limit(6).all()

    center_exams = Exam.query.order_by(Exam.created_at.desc()).limit(8).all()
    exams_data = []
    for exam in center_exams:
        total_subs = ExamSubmission.query.filter_by(exam_id=exam.id).count()
        exams_data.append({
            "exam": exam,
            "total_submissions": total_subs
        })

    # Center activities
    activities = []
    payments = TuitionPayment.query.order_by(TuitionPayment.payment_date.desc()).limit(15).all()
    for p in payments:
        activities.append({
            "type": "payment",
            "title": f"Học phí: {p.student.fullname}",
            "desc": f"Khóa {p.course.course_name}: {p.amount:,} đ ({p.status})",
            "time": dt.datetime.combine(p.payment_date, dt.time.min),
            "icon": "coins"
        })
    teaching_recs = TeachingRecord.query.order_by(TeachingRecord.date.desc()).limit(15).all()
    for tr in teaching_recs:
        activities.append({
            "type": "teaching_record",
            "title": f"Báo giảng: {tr.teacher.fullname}",
            "desc": f"Khóa {tr.course.course_name} ({tr.hours_taught}h) - {'Đã xác nhận' if tr.confirmed_by_teacher else 'Chờ xác nhận'}",
            "time": dt.datetime.combine(tr.date, dt.time.min),
            "icon": "clipboard-check"
        })
    subs = ExamSubmission.query.order_by(ExamSubmission.submitted_at.desc()).limit(15).all()
    for s in subs:
        activities.append({
            "type": "exam",
            "title": f"Bài thi: {s.student.fullname}",
            "desc": f"Môn {s.exam.course.course_name} - {'Đang chờ chấm' if not s.is_graded else f'{s.total_score}/{s.max_score}'}",
            "time": s.submitted_at,
            "icon": "file-question"
        })
    
    activities.sort(key=lambda x: x["time"], reverse=True)
    activities = activities[:10]

    context = {
        "dashboard_scope": "center",
        "total_students": total_students,
        "total_teachers": total_teachers,
        "total_courses": total_courses,
        "total_classes_today": total_classes_today,
        "monthly_revenue": monthly_revenue,
        "unpaid_tuition_count": unpaid_tuition_count,
        "total_teaching_hours": total_teaching_hours,
        "total_teacher_payroll": total_teacher_payroll,
        "timetable_grid": timetable_grid,
        "recent_attendance": recent_attendance,
        "recent_payments": recent_payments,
        "recent_teaching_records": recent_teaching_records,
        "exams_data": exams_data,
        "activities": activities,
        "schedules": all_schedules
    }
    return render_template("index.html", **{**defaults, **context})



@app.route("/courses")
@login_required
def courses():
    query = Course.query
    if current_user.role == "teacher":
        query = query.filter_by(teacher_id=current_user.id)
    elif current_user.role in ("student", "parent"):
        student_ids = get_accessible_student_ids()
        course_ids = [enrollment.course_id for enrollment in Enrollment.query.filter(Enrollment.student_id.in_(student_ids)).all()]
        query = query.filter(Course.id.in_(course_ids)) if course_ids else query.filter(db.text("0=1"))
    return render_template("courses.html", courses=query.order_by(Course.course_name).all())


@app.route("/courses/new", methods=["GET", "POST"])
@login_required
@manager_or_admin_required
def add_course():
    teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all()
    if request.method == "POST":
        course = Course(
            course_name=request.form["course_name"],
            subject=request.form["subject"],
            teacher_id=request.form.get("teacher_id") or None,
            classroom=request.form.get("classroom"),
            description=request.form.get("description"),
            tuition_amount=money(request.form.get("tuition_amount")),
            tuition_type=request.form.get("tuition_type", "monthly"),
            start_date=parse_date(request.form.get("start_date")),
            end_date=parse_date(request.form.get("end_date")),
            is_active=bool(request.form.get("is_active")),
        )
        db.session.add(course)
        db.session.commit()
        flash("Đã tạo khóa học.", "success")
        return redirect(url_for("courses"))
    return render_template("course_form.html", course=None, teachers=teachers)


@app.route("/courses/<int:course_id>/edit", methods=["GET", "POST"])
@login_required
@manager_or_admin_required
def edit_course(course_id):
    course = db.session.get(Course, course_id) or abort(404)
    teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all()
    if request.method == "POST":
        course.course_name = request.form["course_name"]
        course.subject = request.form["subject"]
        course.teacher_id = request.form.get("teacher_id") or None
        course.classroom = request.form.get("classroom")
        course.description = request.form.get("description")
        course.tuition_amount = money(request.form.get("tuition_amount"))
        course.tuition_type = request.form.get("tuition_type", "monthly")
        course.start_date = parse_date(request.form.get("start_date"))
        course.end_date = parse_date(request.form.get("end_date"))
        course.is_active = bool(request.form.get("is_active"))
        db.session.commit()
        flash("Đã cập nhật khóa học.", "success")
        return redirect(url_for("courses"))
    return render_template("course_form.html", course=course, teachers=teachers)


@app.route("/courses/<int:course_id>/delete", methods=["POST"])
@login_required
@manager_or_admin_required
def delete_course(course_id):
    db.session.delete(db.session.get(Course, course_id) or abort(404))
    db.session.commit()
    flash("Đã xóa khóa học.", "success")
    return redirect(url_for("courses"))


@app.route("/courses/<int:course_id>")
@login_required
def course_detail(course_id):
    course = db.session.get(Course, course_id) or abort(404)
    # Check authorization if teacher, student, parent
    if current_user.role == "teacher" and course.teacher_id != current_user.id:
        flash("Bạn không dạy khóa học này.", "danger")
        return redirect(url_for("courses"))
    elif current_user.role in ("student", "parent"):
        student_ids = get_accessible_student_ids()
        is_enrolled = Enrollment.query.filter(Enrollment.student_id.in_(student_ids), Enrollment.course_id == course.id, Enrollment.status == "active").first()
        if not is_enrolled:
            flash("Bạn không có quyền truy cập khóa học này.", "danger")
            return redirect(url_for("courses"))

    # Active enrollments
    active_enrollments = Enrollment.query.filter_by(course_id=course.id, status="active").all()
    # List of all active students for select dropdown (excluding already enrolled)
    enrolled_student_ids = [e.student_id for e in active_enrollments]
    all_active_students = User.query.filter_by(role="student", is_active=True).order_by(User.fullname).all()
    available_students = [s for s in all_active_students if s.id not in enrolled_student_ids]
    
    teachers = User.query.filter_by(role="teacher", is_active=True).order_by(User.fullname).all()
    return render_template(
        "course_detail.html", 
        course=course, 
        enrollments=active_enrollments, 
        available_students=available_students,
        teachers=teachers
    )


# ──────────────────────── SCHEDULES ────────────────────────

@app.route("/schedules", methods=["GET", "POST"])
@login_required
@manager_or_admin_required
def schedules():
    courses_list = Course.query.order_by(Course.course_name).all()
    teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all()
    if request.method == "POST":
        start = parse_time(request.form["start_time"])
        end = parse_time(request.form["end_time"])
        schedule = Schedule(
            course_id=request.form["course_id"],
            teacher_id=request.form["teacher_id"],
            classroom=request.form["classroom"],
            weekday=int(request.form["weekday"]),
            start_time=start,
            end_time=end,
            duration_hours=duration_hours(start, end),
        )
        course = db.session.get(Course, int(request.form["course_id"]))
        if course:
            course.teacher_id = schedule.teacher_id
            course.classroom = schedule.classroom
        db.session.add(schedule)
        db.session.commit()
        flash("Đã tạo lịch học.", "success")
        return redirect(request.form.get("redirect_url") or request.referrer or url_for("schedules"))
    all_schedules = Schedule.query.order_by(Schedule.weekday, Schedule.start_time).all()
    return render_template("schedules.html", schedules=all_schedules, courses=courses_list, teachers=teachers)


@app.route("/schedules/<int:schedule_id>/edit", methods=["GET", "POST"])
@login_required
@manager_or_admin_required
def edit_schedule(schedule_id):
    schedule = db.session.get(Schedule, schedule_id) or abort(404)
    courses_list = Course.query.order_by(Course.course_name).all()
    teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all()
    if request.method == "POST":
        start = parse_time(request.form["start_time"])
        end = parse_time(request.form["end_time"])
        schedule.course_id = request.form["course_id"]
        schedule.teacher_id = request.form["teacher_id"]
        schedule.classroom = request.form["classroom"]
        schedule.weekday = int(request.form["weekday"])
        schedule.start_time = start
        schedule.end_time = end
        schedule.duration_hours = duration_hours(start, end)
        db.session.commit()
        flash("Đã cập nhật lịch học.", "success")
        return redirect(url_for("schedules"))
    return render_template("schedule_form.html", schedule=schedule, courses=courses_list, teachers=teachers)


@app.route("/schedules/<int:schedule_id>/delete", methods=["POST"])
@login_required
@manager_or_admin_required
def delete_schedule(schedule_id):
    db.session.delete(db.session.get(Schedule, schedule_id) or abort(404))
    db.session.commit()
    flash("Đã xóa lịch học.", "success")
    return redirect(url_for("schedules"))


# ──────────────────────── TIMETABLE ────────────────────────

@app.route("/timetable")
@login_required
def center_timetable():
    if current_user.role not in ("admin", "manager"):
        return redirect(url_for("my_schedule"))
    view_mode = request.args.get("view", "week")
    anchor = parse_date(request.args.get("date")) or dt.date.today()
    start, end = month_bounds(anchor.strftime("%Y-%m")) if view_mode == "month" else week_bounds(anchor)
    user_schedules = schedules_for_user(current_user)
    occurrences = schedule_occurrences(user_schedules, start, end)
    timetable_grid = build_timetable_grid(user_schedules, start, end)
    color_map = get_course_color_map(user_schedules)
    return render_template("timetable.html", title="Thời khóa biểu trung tâm", occurrences=occurrences,
                           view_mode=view_mode, anchor=anchor, timetable_grid=timetable_grid, color_map=color_map)


@app.route("/teacher-timetable")
@app.route("/teacher-timetable/<int:teacher_id>")
@login_required
def teacher_timetable(teacher_id=None):
    # Security: teachers can only see their own timetable
    if current_user.role == "teacher":
        teacher_id = current_user.id
    elif current_user.role not in ("admin", "manager"):
        flash("Bạn không có quyền xem lịch giáo viên.", "danger")
        return redirect(url_for("index"))
    teacher = db.session.get(User, teacher_id) if teacher_id else None
    schedules_query = Schedule.query.filter_by(teacher_id=teacher_id) if teacher_id else Schedule.query
    view_mode = request.args.get("view", "week")
    anchor = parse_date(request.args.get("date")) or dt.date.today()
    start, end = month_bounds(anchor.strftime("%Y-%m")) if view_mode == "month" else week_bounds(anchor)
    schedules_list = schedules_query.order_by(Schedule.weekday, Schedule.start_time).all()
    occurrences = schedule_occurrences(schedules_list, start, end)
    timetable_grid = build_timetable_grid(schedules_list, start, end)
    color_map = get_course_color_map(schedules_list)
    return render_template(
        "timetable.html",
        title="Thời khóa biểu giáo viên",
        occurrences=occurrences,
        view_mode=view_mode,
        anchor=anchor,
        selected_user=teacher,
        teachers=User.query.filter_by(role="teacher").all() if current_user.role in ("admin", "manager") else [],
        timetable_grid=timetable_grid,
        color_map=color_map,
    )


@app.route("/student-timetable")
@app.route("/student-timetable/<int:student_id>")
@login_required
def student_timetable(student_id=None):
    # Security: students can only see their own timetable
    if current_user.role == "student":
        student_id = current_user.id
    elif current_user.role == "parent":
        student_id = current_user.linked_student_id
    elif current_user.role not in ("admin", "manager"):
        flash("Bạn không có quyền xem lịch học sinh.", "danger")
        return redirect(url_for("index"))
    # Extra check: if student/parent, don't allow passing arbitrary student_id via URL
    if current_user.role == "student" and student_id != current_user.id:
        student_id = current_user.id
    if current_user.role == "parent" and student_id != current_user.linked_student_id:
        student_id = current_user.linked_student_id

    student = db.session.get(User, student_id) if student_id else None
    if student_id:
        course_ids = [enrollment.course_id for enrollment in Enrollment.query.filter_by(student_id=student_id, status="active").all()]
        schedules_list = Schedule.query.filter(Schedule.course_id.in_(course_ids)).order_by(Schedule.weekday, Schedule.start_time).all() if course_ids else []
    else:
        schedules_list = Schedule.query.order_by(Schedule.weekday, Schedule.start_time).all()
    view_mode = request.args.get("view", "week")
    anchor = parse_date(request.args.get("date")) or dt.date.today()
    start, end = month_bounds(anchor.strftime("%Y-%m")) if view_mode == "month" else week_bounds(anchor)
    timetable_grid = build_timetable_grid(schedules_list, start, end)
    color_map = get_course_color_map(schedules_list)
    return render_template(
        "timetable.html",
        title="Thời khóa biểu học sinh",
        occurrences=schedule_occurrences(schedules_list, start, end),
        view_mode=view_mode,
        anchor=anchor,
        selected_user=student,
        students=User.query.filter_by(role="student").all() if current_user.role in ("admin", "manager") else [],
        timetable_grid=timetable_grid,
        color_map=color_map,
    )


@app.route("/my_schedule")
@login_required
def my_schedule():
    view_mode = request.args.get("view", "week")
    anchor = parse_date(request.args.get("date")) or dt.date.today()
    start, end = month_bounds(anchor.strftime("%Y-%m")) if view_mode == "month" else week_bounds(anchor)
    user_schedules = schedules_for_user(current_user)
    timetable_grid = build_timetable_grid(user_schedules, start, end)
    color_map = get_course_color_map(user_schedules)
    return render_template(
        "timetable.html",
        title="Lịch của tôi",
        occurrences=schedule_occurrences(user_schedules, start, end),
        view_mode=view_mode,
        anchor=anchor,
        timetable_grid=timetable_grid,
        color_map=color_map,
    )


# ──────────────────────── ATTENDANCE ────────────────────────

@app.route("/attendance")
@login_required
@teacher_manager_admin_required
def attendance():
    if current_user.role == "teacher":
        courses_list = Course.query.filter_by(teacher_id=current_user.id).order_by(Course.course_name).all()
    else:
        courses_list = Course.query.order_by(Course.course_name).all()
    records = Attendance.query.order_by(Attendance.date.desc(), Attendance.id.desc()).limit(100).all()
    if current_user.role == "teacher":
        teacher_course_ids = [course.id for course in courses_list]
        records = [record for record in records if record.course_id in teacher_course_ids]
    return render_template("attendance.html", courses=courses_list, records=records)


@app.route("/my-students")
@login_required
def my_students():
    if current_user.role != "teacher":
        flash("Trang này chỉ dành cho giáo viên.", "danger")
        return redirect(url_for("index"))
    students = students_for_teacher(current_user.id)
    return render_template("my_students.html", students=students)


@app.route("/attendance/course/<int:course_id>", methods=["GET", "POST"])
@login_required
@teacher_manager_admin_required
def mark_attendance(course_id):
    course = db.session.get(Course, course_id) or abort(404)
    if current_user.role == "teacher" and course.teacher_id != current_user.id:
        flash("Bạn chỉ được chấm điểm danh lớp mình phụ trách.", "danger")
        return redirect(url_for("attendance"))
    date_value = parse_date(request.form.get("date") or request.args.get("date")) or dt.date.today()
    students = [enrollment.student for enrollment in course.enrollments if enrollment.status == "active"]

    if request.method == "POST":
        for student in students:
            status = request.form.get(f"status_{student.id}", "present")
            record = Attendance.query.filter_by(student_id=student.id, course_id=course.id, date=date_value).first()
            if not record:
                session_schedule = matching_schedule(course, date_value)
                record = Attendance(
                    student_id=student.id,
                    course_id=course.id,
                    date=date_value,
                    schedule_id=session_schedule.id if session_schedule else None,
                )
                db.session.add(record)
            record.status = status
            record.note = request.form.get(f"note_{student.id}")
            record.marked_by_id = current_user.id
        db.session.flush()
        sync_teaching_record(course, date_value)
        db.session.commit()
        flash("Đã lưu điểm danh.", "success")
        return redirect(url_for("mark_attendance", course_id=course.id, date=date_value.isoformat()))

    existing = {record.student_id: record for record in Attendance.query.filter_by(course_id=course.id, date=date_value).all()}
    return render_template("mark_attendance.html", course=course, students=students, date_value=date_value, existing=existing)


# ──────────────────────── PAYROLL ────────────────────────

@app.route("/payroll")
@login_required
def payroll():
    if current_user.role == "teacher":
        teachers = [current_user]
    elif current_user.role in ("admin", "manager"):
        teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all()
    else:
        flash("Bạn không có quyền xem bảng lương.", "danger")
        return redirect(url_for("index"))
    month = request.args.get("month") or month_key()
    start, end = month_bounds(month)
    rows = [{"teacher": teacher, **teaching_summary(teacher.id, start, end)} for teacher in teachers]
    return render_template("payroll.html", rows=rows, month=month)


@app.route("/payroll/<int:teacher_id>")
@login_required
def payroll_detail(teacher_id):
    if current_user.role == "teacher" and teacher_id != current_user.id:
        flash("Bạn chỉ được xem bảng lương của chính mình.", "danger")
        return redirect(url_for("payroll"))
    if current_user.role not in ("admin", "manager", "teacher"):
        flash("Bạn không có quyền xem bảng lương.", "danger")
        return redirect(url_for("index"))

    teacher = db.session.get(User, teacher_id) or abort(404)
    if teacher.role != "teacher":
        abort(404)
    month = request.args.get("month") or month_key()
    start, end = month_bounds(month)
    records = (
        TeachingRecord.query.filter(
            TeachingRecord.teacher_id == teacher.id,
            TeachingRecord.date >= start,
            TeachingRecord.date <= end,
        )
        .order_by(TeachingRecord.date.desc(), TeachingRecord.start_time.desc())
        .all()
    )
    summary = teaching_summary(teacher.id, start, end)
    return render_template("payroll_detail.html", teacher=teacher, records=records, summary=summary, month=month)


# ──────────────────────── TEACHING RECORDS (CONFIRM) ────────────────────────

@app.route("/teaching/my-records")
@login_required
def my_teaching_records():
    if current_user.role not in ("admin", "manager", "teacher"):
        flash("Bạn không có quyền truy cập.", "danger")
        return redirect(url_for("index"))
    if current_user.role == "teacher":
        teacher_id = current_user.id
    else:
        teacher_id = request.args.get("teacher_id", type=int)
    month = request.args.get("month") or month_key()
    start, end = month_bounds(month)
    query = TeachingRecord.query.filter(TeachingRecord.date >= start, TeachingRecord.date <= end)
    if teacher_id:
        query = query.filter_by(teacher_id=teacher_id)
    records = query.order_by(TeachingRecord.date.desc(), TeachingRecord.start_time.desc()).all()
    teachers = User.query.filter_by(role="teacher").order_by(User.fullname).all() if current_user.role in ("admin", "manager") else []
    return render_template("teaching_records.html", records=records, month=month, teachers=teachers, selected_teacher_id=teacher_id)


@app.route("/teaching/confirm/<int:record_id>", methods=["POST"])
@login_required
def confirm_teaching(record_id):
    record = db.session.get(TeachingRecord, record_id) or abort(404)
    if current_user.role == "teacher" and record.teacher_id != current_user.id:
        flash("Bạn chỉ được xác nhận buổi dạy của mình.", "danger")
        return redirect(url_for("my_teaching_records"))
    if current_user.role not in ("admin", "manager", "teacher"):
        abort(403)
    record.confirmed_by_teacher = True
    record.confirmed_at = dt.datetime.utcnow()
    db.session.commit()
    flash("Đã xác nhận buổi dạy.", "success")
    return redirect(request.referrer or url_for("my_teaching_records"))


# ──────────────────────── TUITION ────────────────────────

@app.route("/tuition", methods=["GET", "POST"])
@login_required
def tuition():
    if current_user.role in ("student", "parent"):
        student_ids = get_accessible_student_ids()
        payments = TuitionPayment.query.filter(TuitionPayment.student_id.in_(student_ids)).order_by(TuitionPayment.payment_date.desc()).all()
        return render_template("tuition.html", payments=payments, students=[], courses=[], readonly=True)

    if current_user.role not in ("admin", "manager"):
        flash("Bạn không có quyền quản lý học phí.", "danger")
        return redirect(url_for("index"))

    if request.method == "POST":
        payment = TuitionPayment(
            student_id=request.form["student_id"],
            course_id=request.form["course_id"],
            amount=money(request.form.get("amount")),
            payment_type=request.form.get("payment_type", "monthly"),
            payment_date=parse_date(request.form.get("payment_date")) or dt.date.today(),
            status=request.form.get("status", "paid"),
            note=request.form.get("note"),
        )
        db.session.add(payment)
        db.session.commit()
        flash("Đã ghi nhận học phí.", "success")
        return redirect(url_for("tuition"))

    payments = TuitionPayment.query.order_by(TuitionPayment.payment_date.desc()).all()
    return render_template(
        "tuition.html",
        payments=payments,
        students=User.query.filter_by(role="student").order_by(User.fullname).all(),
        courses=Course.query.order_by(Course.course_name).all(),
        readonly=False,
    )


@app.route("/tuition/<int:payment_id>/edit", methods=["POST"])
@login_required
@manager_or_admin_required
def edit_tuition(payment_id):
    payment = db.session.get(TuitionPayment, payment_id) or abort(404)
    payment.amount = money(request.form.get("amount"))
    payment.status = request.form.get("status", payment.status)
    payment.note = request.form.get("note")
    db.session.commit()
    flash("Đã cập nhật học phí.", "success")
    return redirect(url_for("tuition"))


@app.route("/users/<role_name>")
@login_required
@manager_or_admin_required
def manage_users(role_name):
    if role_name not in ("students", "teachers", "parents", "staff"):
        abort(404)
    role_map = {"students": ["student"], "teachers": ["teacher"], "parents": ["parent"], "staff": ["admin", "manager"]}
    users = User.query.filter(User.role.in_(role_map[role_name])).order_by(User.fullname).all()
    students = User.query.filter_by(role="student").order_by(User.fullname).all()
    return render_template("manage_users.html", users=users, role_name=role_name, students=students)


@app.route("/users", methods=["POST"])
@login_required
@manager_or_admin_required
def add_user():
    role = request.form["role"]
    if current_user.role == "manager" and role == "admin":
        flash("Quản lý không được tạo tài khoản admin.", "danger")
        return redirect(request.referrer or url_for("index"))
    user = User(
        username=request.form["username"],
        fullname=request.form["fullname"],
        user_code=request.form["user_code"],
        role=role,
        phone=request.form.get("phone"),
        position=request.form.get("position"),
        salary_rate_per_hour=money(request.form.get("salary_rate_per_hour")),
        linked_student_id=request.form.get("linked_student_id") or None,
        is_active=True,
    )
    user.set_password(request.form.get("password") or "123456")
    db.session.add(user)
    db.session.commit()
    flash("Đã thêm người dùng.", "success")
    return redirect(request.referrer or url_for("manage_users", role_name="students"))


@app.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
@login_required
@manager_or_admin_required
def edit_user(user_id):
    user = db.session.get(User, user_id) or abort(404)
    if current_user.role == "manager" and user.role == "admin":
        flash("Quản lý không được thay đổi tài khoản admin.", "danger")
        return redirect(url_for("manage_users", role_name="staff"))
    students = User.query.filter_by(role="student").order_by(User.fullname).all()
    if request.method == "POST":
        new_role = request.form["role"]
        if current_user.role == "manager" and (user.role == "admin" or new_role == "admin"):
            flash("Quản lý không được thay đổi quyền admin.", "danger")
            return redirect(url_for("manage_users", role_name="staff"))
        user.fullname = request.form["fullname"]
        user.username = request.form["username"]
        user.user_code = request.form["user_code"]
        user.role = new_role
        user.phone = request.form.get("phone")
        user.position = request.form.get("position")
        user.salary_rate_per_hour = money(request.form.get("salary_rate_per_hour"))
        user.linked_student_id = request.form.get("linked_student_id") or None
        user.is_active = bool(request.form.get("is_active"))
        db.session.commit()
        flash("Đã cập nhật người dùng.", "success")
        return redirect(url_for("manage_users", role_name="teachers" if user.role == "teacher" else "students"))
    return render_template("edit_user.html", user=user, students=students)


@app.route("/users/<int:user_id>/delete", methods=["POST"])
@login_required
@manager_or_admin_required
def delete_user(user_id):
    user = db.session.get(User, user_id) or abort(404)
    if user.id == current_user.id or (current_user.role == "manager" and user.role == "admin"):
        flash("Không thể xóa tài khoản này.", "danger")
        return redirect(request.referrer or url_for("index"))
    db.session.delete(user)
    db.session.commit()
    flash("Đã xóa người dùng.", "success")
    return redirect(request.referrer or url_for("index"))


@app.route("/users/<int:user_id>/reset-password", methods=["POST"])
@login_required
@manager_or_admin_required
def reset_password(user_id):
    user = db.session.get(User, user_id) or abort(404)
    if current_user.role == "manager" and user.role == "admin":
        flash("Quản lý không được reset mật khẩu admin.", "danger")
        return redirect(request.referrer or url_for("index"))
    user.set_password(request.form.get("password") or "123456")
    db.session.commit()
    flash("Đã reset mật khẩu.", "success")
    return redirect(request.referrer or url_for("index"))


@app.route("/enrollments", methods=["POST"])
@login_required
@manager_or_admin_required
def add_enrollment():
    student_id_raw = request.form["student_id"]
    course_id = int(request.form["course_id"])
    
    student_id = resolve_student_id(student_id_raw)
    if not student_id:
        flash(f"Không tìm thấy học sinh với mã hoặc tên đăng nhập '{student_id_raw}'.", "danger")
        return redirect(request.referrer or url_for("courses"))
        
    existing = Enrollment.query.filter_by(student_id=student_id, course_id=course_id).first()
    if existing:
        if existing.status == "active":
            flash("Học sinh này đã ghi danh vào khóa học rồi.", "warning")
        else:
            existing.status = "active"
            db.session.commit()
            flash("Đã kích hoạt lại ghi danh cho học sinh này.", "success")
        return redirect(request.referrer or url_for("courses"))
        
    enrollment = Enrollment(student_id=student_id, course_id=course_id, status="active")
    db.session.add(enrollment)
    db.session.commit()
    flash("Đã ghi danh học sinh thành công.", "success")
    return redirect(request.referrer or url_for("courses"))


# ──────────────────────── QUESTION BANK & IMPORT ────────────────────────

@app.route("/question-banks")
@login_required
@teacher_manager_admin_required
def question_bank_list():
    if current_user.role in ("admin", "manager"):
        banks = QuestionBank.query.order_by(QuestionBank.updated_at.desc()).all()
    else:
        banks = QuestionBank.query.filter_by(created_by=current_user.id).order_by(QuestionBank.updated_at.desc()).all()
    return render_template("question_banks.html", banks=banks)


@app.route("/question-banks/new", methods=["GET", "POST"])
@login_required
@teacher_manager_admin_required
def create_question_bank():
    if request.method == "POST":
        bank = QuestionBank(
            name=request.form["name"].strip(),
            subject=request.form["subject"].strip(),
            grade=request.form.get("grade", "").strip(),
            topic=request.form.get("topic", "").strip(),
            description=request.form.get("description", "").strip(),
            created_by=current_user.id
        )
        db.session.add(bank)
        db.session.commit()
        flash("Đã tạo ngân hàng câu hỏi thành công.", "success")
        return redirect(url_for("question_bank_detail", bank_id=bank.id))
    return render_template("question_bank_form.html", bank=None)


@app.route("/question-banks/<int:bank_id>")
@login_required
@teacher_manager_admin_required
def question_bank_detail(bank_id):
    bank = db.session.get(QuestionBank, bank_id) or abort(404)
    if current_user.role == "teacher" and bank.created_by != current_user.id:
        abort(403)
        
    q_type_filter = request.args.get("type", "")
    diff_filter = request.args.get("difficulty", "")
    search_q = request.args.get("q", "").strip()
    
    query = BankQuestion.query.filter_by(bank_id=bank.id)
    if q_type_filter:
        query = query.filter_by(question_type=q_type_filter)
    if diff_filter:
        query = query.filter_by(difficulty_level=diff_filter)
    if search_q:
        query = query.filter(BankQuestion.question_text.ilike(f"%{search_q}%"))
        
    questions = query.order_by(BankQuestion.id.desc()).all()
    questions_json = [q.to_dict() for q in questions]
    return render_template("question_bank_detail.html", bank=bank, questions=questions, questions_json=questions_json)


@app.route("/question-banks/<int:bank_id>/edit", methods=["GET", "POST"])
@login_required
@teacher_manager_admin_required
def edit_question_bank(bank_id):
    bank = db.session.get(QuestionBank, bank_id) or abort(404)
    if current_user.role == "teacher" and bank.created_by != current_user.id:
        abort(403)
        
    if request.method == "POST":
        bank.name = request.form["name"].strip()
        bank.subject = request.form["subject"].strip()
        bank.grade = request.form.get("grade", "").strip()
        bank.topic = request.form.get("topic", "").strip()
        bank.description = request.form.get("description", "").strip()
        bank.updated_at = dt.datetime.utcnow()
        db.session.commit()
        flash("Đã cập nhật thông tin ngân hàng câu hỏi.", "success")
        return redirect(url_for("question_bank_detail", bank_id=bank.id))
    return render_template("question_bank_form.html", bank=bank)


@app.route("/question-banks/<int:bank_id>/delete", methods=["POST"])
@login_required
@teacher_manager_admin_required
def delete_question_bank(bank_id):
    bank = db.session.get(QuestionBank, bank_id) or abort(404)
    if current_user.role == "teacher" and bank.created_by != current_user.id:
        abort(403)

    try:
        # 1. Clean up linked ImportSessions & their temporary files
        sessions = ImportSession.query.filter_by(bank_id=bank.id).all()
        for session in sessions:
            ImportService.cleanup_temp_file(session)
            db.session.delete(session)

        # 2. Delete all question images and option images on disk
        for q in bank.questions:
            if q.image_url:
                rel_path = q.image_url.lstrip("/")
                full_path = os.path.join(app.root_path, rel_path)
                if os.path.exists(full_path) and os.path.isfile(full_path):
                    try:
                        os.remove(full_path)
                    except OSError:
                        pass
            for opt in q.options:
                if hasattr(opt, "image_url") and opt.image_url:
                    opt_rel = opt.image_url.lstrip("/")
                    opt_full = os.path.join(app.root_path, opt_rel)
                    if os.path.exists(opt_full) and os.path.isfile(opt_full):
                        try:
                            os.remove(opt_full)
                        except OSError:
                            pass

        # 3. Delete any import_temp folder associated with this bank
        bank_temp_dir = os.path.join(app.root_path, "static", "uploads", "import_temp", str(bank.id))
        if os.path.exists(bank_temp_dir):
            import shutil
            shutil.rmtree(bank_temp_dir, ignore_errors=True)

        # 4. Unlink foreign keys in Exam & ExamQuestion
        exams = Exam.query.filter_by(question_bank_id=bank.id).all()
        for ex in exams:
            ex.question_bank_id = None

        q_ids = [q.id for q in bank.questions]
        if q_ids:
            exam_qs = ExamQuestion.query.filter(ExamQuestion.bank_question_id.in_(q_ids)).all()
            for eq in exam_qs:
                eq.bank_question_id = None

        # 5. Delete the QuestionBank record (cascades to questions & options)
        db.session.delete(bank)
        db.session.commit()

        # 6. Purge any orphan unreferenced image files from static/uploads/questions
        ImportService.cleanup_orphan_question_images()

        flash("Đã xóa hoàn toàn ngân hàng câu hỏi và toàn bộ dữ liệu liên quan.", "success")
    except Exception as e:
        db.session.rollback()
        flash(f"Lỗi khi xóa ngân hàng câu hỏi: {str(e)}", "danger")

    return redirect(url_for("question_bank_list"))


@app.route("/question-banks/<int:bank_id>/questions/add", methods=["POST"])
@login_required
@teacher_manager_admin_required
def add_bank_question(bank_id):
    bank = db.session.get(QuestionBank, bank_id) or abort(404)
    if current_user.role == "teacher" and bank.created_by != current_user.id:
        abort(403)
        
    context = request.form.get("context", "").strip()
    question_text = request.form.get("question_text", "").strip()
    question_type = request.form.get("question_type", "single")
    if question_type == "truefalse":
        question_type = "true_false"
    difficulty_level = request.form.get("difficulty_level", "nhan_biet")
    points = float(request.form.get("points", 1.0))
    explanation = request.form.get("explanation", "").strip()
    
    if not question_text:
        flash("Nội dung câu hỏi không được để trống.", "danger")
        return redirect(url_for("question_bank_detail", bank_id=bank.id))
        
    bq = BankQuestion(
        bank_id=bank.id,
        question_text=question_text,
        context=context,
        question_type=question_type,
        difficulty_level=difficulty_level,
        points=points,
        explanation=explanation,
        subject=bank.subject,
        grade=bank.grade,
        topic=bank.topic,
        created_by=current_user.id
    )
    db.session.add(bq)
    db.session.flush()
    
    # Process options based on question_type
    if question_type == "single":
        option_texts = request.form.getlist("option_text")
        correct_index = int(request.form.get("correct_option", 0)) if request.form.get("correct_option") else 0
        for idx, opt_text in enumerate(option_texts):
            if opt_text.strip():
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=opt_text.strip(),
                    is_correct=(idx == correct_index),
                    order_index=idx
                )
                db.session.add(opt)
    elif question_type == "multiple_choice":
        option_texts = request.form.getlist("option_text")
        correct_indices = [int(x) for x in request.form.getlist("correct_option") if x.isdigit()]
        for idx, opt_text in enumerate(option_texts):
            if opt_text.strip():
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=opt_text.strip(),
                    is_correct=(idx in correct_indices),
                    order_index=idx
                )
                db.session.add(opt)
    elif question_type == "true_false":
        option_texts = request.form.getlist("option_text")
        correct_values = request.form.getlist("correct_option")
        labels = ["a", "b", "c", "d", "e", "f"]
        for idx, opt_text in enumerate(option_texts):
            if opt_text.strip():
                stmt_label = labels[idx] if idx < len(labels) else f"s{idx+1}"
                full_text = f"{stmt_label}) {opt_text.strip()}" if not opt_text.strip().startswith(f"{stmt_label})") else opt_text.strip()
                is_corr = bool(correct_values[idx]) if idx < len(correct_values) and correct_values[idx] != "" else False
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=full_text,
                    is_correct=is_corr,
                    order_index=idx
                )
                db.session.add(opt)
    elif question_type == "short_answer":
        answers_raw = request.form.get("short_answers", "").strip()
        answers = [a.strip() for a in answers_raw.split("|") if a.strip()]
        for idx, ans in enumerate(answers):
            opt = BankQuestionOption(
                question_id=bq.id,
                option_text=ans,
                is_correct=True,
                order_index=idx
            )
            db.session.add(opt)
            
    db.session.commit()
    flash("Đã thêm câu hỏi mới vào ngân hàng.", "success")
    return redirect(url_for("question_bank_detail", bank_id=bank.id))


@app.route("/questions/<int:question_id>/delete", methods=["POST"])
@login_required
@teacher_manager_admin_required
def delete_bank_question(question_id):
    bq = db.session.get(BankQuestion, question_id) or abort(404)
    if current_user.role == "teacher" and bq.bank.created_by != current_user.id:
        abort(403)
    bank_id = bq.bank_id
    db.session.delete(bq)
    db.session.commit()
    flash("Đã xóa câu hỏi.", "success")
    return redirect(url_for("question_bank_detail", bank_id=bank_id))


@app.route("/questions/<int:question_id>/update", methods=["POST"])
@login_required
@teacher_manager_admin_required
def update_bank_question_json(question_id):
    bq = db.session.get(BankQuestion, question_id)
    if not bq:
        return jsonify({"success": False, "message": "Không tìm thấy câu hỏi."}), 404
    if current_user.role == "teacher" and bq.bank.created_by != current_user.id:
        return jsonify({"success": False, "message": "Không có quyền chỉnh sửa."}), 403

    data = request.get_json() or {}
    bq.question_text = data.get("question_text", bq.question_text).strip()
    bq.context = data.get("context", bq.context or "").strip()
    q_type = data.get("question_type", bq.question_type)
    if q_type == "truefalse":
        q_type = "true_false"
    bq.question_type = q_type
    bq.difficulty_level = data.get("difficulty_level", bq.difficulty_level)
    bq.explanation = data.get("explanation", bq.explanation or "").strip()
    bq.image_url = data.get("image_url", bq.image_url or "").strip()

    # Clear existing options
    BankQuestionOption.query.filter_by(question_id=bq.id).delete()

    # Re-create options based on question_type
    if q_type in ("single", "multiple_choice"):
        opts_data = data.get("options", [])
        for idx, opt_item in enumerate(opts_data):
            opt_text = opt_item.get("text") or opt_item.get("option_text") or ""
            if opt_text.strip():
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=opt_text.strip(),
                    is_correct=bool(opt_item.get("is_correct")),
                    order_index=idx
                )
                db.session.add(opt)
    elif q_type == "true_false":
        stmts_data = data.get("statements", [])
        labels = ["a", "b", "c", "d", "e", "f"]
        for idx, stmt in enumerate(stmts_data):
            s_text = stmt.get("text", "").strip()
            s_id = stmt.get("id") or (labels[idx] if idx < len(labels) else f"s{idx+1}")
            if s_text:
                full_text = f"{s_id}) {s_text}" if not s_text.startswith(f"{s_id})") else s_text
                is_corr = bool(stmt.get("answer")) if stmt.get("answer") is not None else False
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=full_text,
                    is_correct=is_corr,
                    order_index=idx
                )
                db.session.add(opt)
    elif q_type == "short_answer":
        opts_data = data.get("options", [])
        for idx, opt_item in enumerate(opts_data):
            opt_text = opt_item.get("text") or opt_item.get("option_text") or ""
            if opt_text.strip():
                opt = BankQuestionOption(
                    question_id=bq.id,
                    option_text=opt_text.strip(),
                    is_correct=True,
                    order_index=idx
                )
                db.session.add(opt)

    db.session.commit()
    return jsonify({"success": True, "question": bq.to_dict()})


@app.route("/questions/<int:question_id>/toggle-answer", methods=["POST"])
@login_required
@teacher_manager_admin_required
def toggle_bank_question_answer_json(question_id):
    bq = db.session.get(BankQuestion, question_id)
    if not bq:
        return jsonify({"success": False, "message": "Không tìm thấy câu hỏi."}), 404
    if current_user.role == "teacher" and bq.bank.created_by != current_user.id:
        return jsonify({"success": False, "message": "Không có quyền chỉnh sửa."}), 403

    data = request.get_json() or {}
    option_idx = data.get("option_index")
    is_correct = data.get("is_correct")
    stmt_id = data.get("statement_id")

    opts = BankQuestionOption.query.filter_by(question_id=bq.id).order_by(BankQuestionOption.order_index).all()
    
    if bq.question_type == "single":
        if option_idx is not None and 0 <= option_idx < len(opts):
            for idx, opt in enumerate(opts):
                opt.is_correct = (idx == option_idx)
    elif bq.question_type == "multiple_choice":
        if option_idx is not None and 0 <= option_idx < len(opts):
            if is_correct is not None:
                opts[option_idx].is_correct = bool(is_correct)
            else:
                opts[option_idx].is_correct = not opts[option_idx].is_correct
    elif bq.question_type in ("true_false", "truefalse"):
        if option_idx is not None and 0 <= option_idx < len(opts):
            if is_correct is not None:
                opts[option_idx].is_correct = bool(is_correct)
            else:
                opts[option_idx].is_correct = not opts[option_idx].is_correct
        elif stmt_id:
            for opt in opts:
                if opt.option_text.lower().startswith(stmt_id.lower() + ")") or opt.option_text.lower().startswith(stmt_id.lower() + "."):
                    opt.is_correct = bool(is_correct)
                    break

    db.session.commit()
    return jsonify({"success": True, "question": bq.to_dict()})


# ──────────── IMPORT WIZARD & API ROUTES ────────────

@app.route("/question-banks/import")
@login_required
@teacher_manager_admin_required
def import_questions_page():
    if current_user.role in ("admin", "manager"):
        banks = QuestionBank.query.order_by(QuestionBank.name).all()
    else:
        banks = QuestionBank.query.filter_by(created_by=current_user.id).order_by(QuestionBank.name).all()
    return render_template("question_bank_import.html", banks=banks)


@app.route("/question-banks/import/upload", methods=["POST"])
@login_required
@teacher_manager_admin_required
def import_upload():
    if "file" not in request.files:
        return jsonify({"success": False, "message": "Không tìm thấy file upload."}), 400
    file = request.files["file"]
    if not file or not file.filename:
        return jsonify({"success": False, "message": "File không hợp lệ."}), 400
        
    bank_id = request.form.get("bank_id")
    bank_id = int(bank_id) if bank_id and bank_id.isdigit() else None
    
    try:
        session = ImportService.create_and_process_session(
            user_id=current_user.id,
            file_obj=file,
            filename=file.filename,
            bank_id=bank_id
        )
        return jsonify({
            "success": True,
            "session_id": session.id,
            "questions": session.get_parsed_questions(),
            "stats": {
                "total": session.total_count,
                "valid": session.valid_count,
                "warning": session.warning_count,
                "error": session.error_count
            }
        })
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 400


@app.route("/question-banks/import/<session_id>/preview", methods=["GET"])
@login_required
@teacher_manager_admin_required
def import_preview_api(session_id):
    session = ImportService.get_session(session_id)
    if not session:
        return jsonify({"success": False, "message": "Không tìm thấy phiên import."}), 404
    return jsonify({
        "success": True,
        "session_id": session.id,
        "status": session.status,
        "questions": session.get_parsed_questions(),
        "stats": {
            "total": session.total_count,
            "valid": session.valid_count,
            "warning": session.warning_count,
            "error": session.error_count
        }
    })


@app.route("/question-banks/import/<session_id>/update", methods=["POST"])
@login_required
@teacher_manager_admin_required
def import_update_api(session_id):
    data = request.json or {}
    q_index = data.get("index")
    question_data = data.get("question")
    if q_index is None or not question_data:
        return jsonify({"success": False, "message": "Dữ liệu không hợp lệ."}), 400

    try:
        session = ImportService.update_preview_question(session_id, q_index, question_data)
        return jsonify({
            "success": True,
            "questions": session.get_parsed_questions(),
            "stats": {
                "total": session.total_count,
                "valid": session.valid_count,
                "warning": session.warning_count,
                "error": session.error_count
            }
        })
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 400


@app.route("/question-banks/import/<session_id>/confirm", methods=["POST"])
@login_required
@teacher_manager_admin_required
def import_confirm_api(session_id):
    data = request.json or {}
    bank_id = data.get("bank_id")
    new_bank_name = data.get("new_bank_name")
    subject = data.get("subject", "Khác")
    grade = data.get("grade", "")

    if not bank_id and not new_bank_name:
        return jsonify({"success": False, "message": "Vui lòng chọn hoặc nhập tên ngân hàng câu hỏi mới."}), 400

    if not bank_id and new_bank_name:
        bank = QuestionBank(
            name=new_bank_name.strip(),
            subject=subject,
            grade=grade,
            created_by=current_user.id
        )
        db.session.add(bank)
        db.session.commit()
        bank_id = bank.id
    else:
        bank_id = int(bank_id)

    try:
        res = ImportService.confirm_import(session_id, bank_id, current_user.id)
        return jsonify(res)
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 400


@app.route("/question-banks/import/<session_id>/cancel", methods=["POST"])
@login_required
@teacher_manager_admin_required
def import_cancel_api(session_id):
    ImportService.cancel_session(session_id)
    return jsonify({"success": True, "message": "Đã hủy phiên import."})


@app.route("/question-banks/import/template/docx")
@login_required
@teacher_manager_admin_required
def download_docx_template():
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("NGÂN HÀNG CÂU HỎI MẪU - EDUCENTER", level=1)
        doc.add_paragraph("Hướng dẫn: Điền câu hỏi theo mẫu bên dưới. Hỗ trợ 5 loại câu hỏi: Trắc nghiệm (1 đáp án), Nhiều đáp án, Đúng/Sai THPT, Trả lời ngắn, Tự luận.")
        doc.add_paragraph("")

        # 1. Single Choice
        doc.add_paragraph("CÂU 1. Cho hàm số f(x) = x² - 2x + 3. Giá trị f(1) bằng bao nhiêu?")
        doc.add_paragraph("A. 1")
        doc.add_paragraph("B. 2")
        doc.add_paragraph("C. 3")
        doc.add_paragraph("D. 4")
        doc.add_paragraph("Đáp án: B")
        doc.add_paragraph("Loại: Trắc nghiệm")
        doc.add_paragraph("Mức độ: Nhận biết")
        doc.add_paragraph("Giải thích: Thay x = 1 vào hàm số ta được f(1) = 1 - 2 + 3 = 2.")
        doc.add_paragraph("")

        # 2. Multiple Choice
        doc.add_paragraph("CÂU 2. Cho a > 0. Các mệnh đề nào sau đây đúng?")
        doc.add_paragraph("A. a² > 0")
        doc.add_paragraph("B. √a > 0")
        doc.add_paragraph("C. a + 1 > 0")
        doc.add_paragraph("D. a² < 0")
        doc.add_paragraph("Đáp án: A, B, C")
        doc.add_paragraph("Loại: Trắc nghiệm nhiều đáp án")
        doc.add_paragraph("Mức độ: Vận dụng")
        doc.add_paragraph("")

        # 3. True/False Group THPT
        doc.add_paragraph("CÂU 3. Cho hàm số f(x) = x³ - 3x. Xét tính đúng sai của các mệnh đề sau:")
        doc.add_paragraph("a) f(0) = 0. (Đúng)")
        doc.add_paragraph("b) Hàm số đồng biến trên toàn bộ R. (Sai)")
        doc.add_paragraph("c) Phương trình f(x) = 0 có ba nghiệm thực phân biệt. (Đúng)")
        doc.add_paragraph("d) Đồ thị hàm số đi qua điểm M(1; -2). (Đúng)")
        doc.add_paragraph("Loại: Đúng/Sai THPT")
        doc.add_paragraph("Mức độ: Vận dụng")
        doc.add_paragraph("")

        # 4. Short Answer
        doc.add_paragraph("CÂU 4. Tính diện tích hình vuông có độ dài cạnh bằng 5 cm.")
        doc.add_paragraph("Đáp án: 25 | 25 cm2")
        doc.add_paragraph("Loại: Trả lời ngắn")
        doc.add_paragraph("Mức độ: Nhận biết")
        doc.add_paragraph("")

        # 5. Essay
        doc.add_paragraph("CÂU 5. Cho hình chóp S.ABCD có đáy ABCD là hình vuông cạnh a. Chứng minh rằng (SAB) ⊥ (ABCD).")
        doc.add_paragraph("Loại: Tự luận")
        doc.add_paragraph("Mức độ: Vận dụng cao")

        from io import BytesIO
        target_stream = BytesIO()
        doc.save(target_stream)
        target_stream.seek(0)
        from flask import send_file
        return send_file(target_stream, as_attachment=True, download_name="Mau_Cau_Hoi_EduCenter.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    except Exception as e:
        flash(f"Lỗi tạo template DOCX: {e}", "danger")
        return redirect(url_for("import_questions_page"))


@app.route("/question-banks/import/template/csv")
@login_required
@teacher_manager_admin_required
def download_csv_template():
    csv_data = (
        "question_text,question_type,difficulty_level,option_a,option_b,option_c,option_d,correct_answer,explanation\n"
        '"Thủ đô của Việt Nam là gì?",single,nhan_biet,"Hà Nội","HCM","Đà Nẵng","Cần Thơ",A,"Hà Nội là thủ đô"\n'
        '"Trái đất quay quanh mặt trời",truefalse,nhan_biet,"Đúng","Sai","","",A,"Định luật Kepler"\n'
        '"Ký hiệu hóa học của Hydro?",short_answer,van_dung,"","","","","H|Hydro","Hydro có số hiệu nguyên tử 1"\n'
    )
    from flask import Response
    return Response(
        csv_data,
        mimetype="text/csv; charset=utf-8-sig",
        headers={"Content-Disposition": "attachment; filename=Mau_Cau_Hoi_EduCenter.csv"}
    )


# ──────────────────────── EXAMS ────────────────────────

@app.route("/exams")
@login_required
def exam_list():
    if current_user.role in ("admin", "manager"):
        exams = Exam.query.order_by(Exam.created_at.desc()).all()
    elif current_user.role == "teacher":
        exams = Exam.query.filter_by(created_by=current_user.id).order_by(Exam.created_at.desc()).all()
    elif current_user.role in ("student", "parent"):
        student_id = current_user.id if current_user.role == "student" else current_user.linked_student_id
        if student_id:
            course_ids = [e.course_id for e in Enrollment.query.filter_by(student_id=student_id, status="active").all()]
            exams = Exam.query.filter(Exam.course_id.in_(course_ids), Exam.is_active == True).order_by(Exam.created_at.desc()).all() if course_ids else []
        else:
            exams = []
    else:
        exams = []
    submissions_map = {}
    if current_user.role == "student":
        for sub in ExamSubmission.query.filter_by(student_id=current_user.id).all():
            submissions_map[sub.exam_id] = sub
    return render_template("exam_list.html", exams=exams, submissions_map=submissions_map)


@app.route("/exams/new", methods=["GET", "POST"])
@login_required
def create_exam():
    if current_user.role not in ("admin", "manager", "teacher"):
        abort(403)
    if current_user.role == "teacher":
        courses_list = Course.query.filter_by(teacher_id=current_user.id).order_by(Course.course_name).all()
        banks_list = QuestionBank.query.filter_by(created_by=current_user.id).order_by(QuestionBank.name).all()
    else:
        courses_list = Course.query.order_by(Course.course_name).all()
        banks_list = QuestionBank.query.order_by(QuestionBank.name).all()

    if request.method == "POST":
        creation_mode = request.form.get("creation_mode", "manual")

        course_id_raw = request.form.get("course_id")
        if not course_id_raw:
            flash("Vui lòng chọn lớp học / khóa học.", "danger")
            return redirect(url_for("create_exam"))

        exam = Exam(
            title=request.form["title"],
            exam_category=request.form.get("exam_category", "Kiểm tra thường kỳ"),
            course_id=int(course_id_raw),
            created_by=current_user.id,
            duration_minutes=int(request.form.get("duration_minutes", 30)),
            shuffle_questions=bool(request.form.get("shuffle_questions")),
            shuffle_answers=bool(request.form.get("shuffle_answers")),
            is_active=bool(request.form.get("is_active", True)),
        )
        start_time_str = request.form.get("start_time")
        end_time_str = request.form.get("end_time")
        if start_time_str:
            exam.start_time = dt.datetime.strptime(start_time_str, "%Y-%m-%dT%H:%M")
        if end_time_str:
            exam.end_time = dt.datetime.strptime(end_time_str, "%Y-%m-%dT%H:%M")

        # Verify teacher owns the course
        if current_user.role == "teacher":
            course = db.session.get(Course, exam.course_id)
            if not course or course.teacher_id != current_user.id:
                flash("Bạn chỉ được tạo bài kiểm tra cho lớp mình.", "danger")
                return redirect(url_for("exam_list"))

        if creation_mode == "bank":
            bank_id = int(request.form["question_bank_id"])
            bank = db.session.get(QuestionBank, bank_id)
            if not bank:
                flash("Ngân hàng câu hỏi không tồn tại.", "danger")
                return redirect(url_for("create_exam"))

            exam.question_bank_id = bank.id
            exam.total_questions = int(request.form.get("total_questions", 10))

            diff_cfg = {
                "nhan_biet": int(request.form.get("diff_nhan_biet", 40)),
                "van_dung": int(request.form.get("diff_van_dung", 30)),
                "van_dung_cao": int(request.form.get("diff_van_dung_cao", 30))
            }
            type_cfg = {
                "single": int(request.form.get("type_single", 50)),
                "truefalse": int(request.form.get("type_truefalse", 30)),
                "short_answer": int(request.form.get("type_short_answer", 20))
            }
            exam.difficulty_config = json.dumps(diff_cfg)
            exam.question_type_config = json.dumps(type_cfg)

            db.session.add(exam)
            db.session.flush()

            selected_bqs = _select_random_bank_questions(bank.id, exam.total_questions, diff_cfg, type_cfg)

            for idx, bq in enumerate(selected_bqs):
                eq = ExamQuestion(
                    exam_id=exam.id,
                    bank_question_id=bq.id,
                    question_text=bq.question_text,
                    question_type=bq.question_type,
                    difficulty_level=bq.difficulty_level,
                    explanation=bq.explanation,
                    points=bq.points,
                    order_index=idx
                )
                db.session.add(eq)
                db.session.flush()

                for opt in bq.options:
                    ea = ExamAnswer(
                        question_id=eq.id,
                        answer_text=opt.option_text,
                        is_correct=opt.is_correct
                    )
                    db.session.add(ea)
        else:
            db.session.add(exam)
            db.session.flush()

            q_idx = 0
            while f"question_{q_idx}_text" in request.form:
                q_text = request.form[f"question_{q_idx}_text"]
                q_type = request.form.get(f"question_{q_idx}_type", "single")
                q_points = float(request.form.get(f"question_{q_idx}_points", 1))
                if q_text.strip():
                    question = ExamQuestion(
                        exam_id=exam.id,
                        question_text=q_text,
                        question_type=q_type,
                        points=q_points,
                        order_index=q_idx,
                    )
                    db.session.add(question)
                    db.session.flush()

                    a_idx = 0
                    while f"question_{q_idx}_answer_{a_idx}_text" in request.form:
                        a_text = request.form[f"question_{q_idx}_answer_{a_idx}_text"]
                        if a_text.strip():
                            if q_type == "truefalse":
                                is_correct = request.form.get(f"question_{q_idx}_correct") == str(a_idx)
                            elif q_type == "multi":
                                is_correct = request.form.get(f"question_{q_idx}_answer_{a_idx}_correct") == "1"
                            else:
                                is_correct = request.form.get(f"question_{q_idx}_correct") == str(a_idx)
                            answer = ExamAnswer(
                                question_id=question.id,
                                answer_text=a_text,
                                is_correct=is_correct,
                            )
                            db.session.add(answer)
                        a_idx += 1
                    q_idx += 1

        db.session.commit()
        flash("Đã tạo bài kiểm tra mới thành công.", "success")
        return redirect(url_for("exam_list"))

    return render_template("exam_form.html", exam=None, courses=courses_list, banks=banks_list)


def _select_random_bank_questions(bank_id, total_needed, diff_cfg, type_cfg):
    """Selection algorithm for building random exams from QuestionBank."""
    all_bqs = BankQuestion.query.filter_by(bank_id=bank_id).all()
    if not all_bqs:
        return []
    if len(all_bqs) <= total_needed:
        random.shuffle(all_bqs)
        return all_bqs

    selected = []
    remaining_bqs = list(all_bqs)

    diff_targets = {
        "nhan_biet": round(total_needed * (diff_cfg.get("nhan_biet", 40) / 100)),
        "van_dung": round(total_needed * (diff_cfg.get("van_dung", 30) / 100)),
        "van_dung_cao": round(total_needed * (diff_cfg.get("van_dung_cao", 30) / 100)),
    }

    diff_sum = sum(diff_targets.values())
    if diff_sum != total_needed:
        diff_targets["nhan_biet"] += (total_needed - diff_sum)

    for diff, target_count in diff_targets.items():
        candidates = [q for q in remaining_bqs if q.difficulty_level == diff]
        picked = random.sample(candidates, min(len(candidates), target_count))
        selected.extend(picked)
        for q in picked:
            remaining_bqs.remove(q)

    if len(selected) < total_needed and remaining_bqs:
        needed = total_needed - len(selected)
        picked = random.sample(remaining_bqs, min(len(remaining_bqs), needed))
        selected.extend(picked)

    random.shuffle(selected)
    return selected


@app.route("/api/exam/check-bank", methods=["POST"])
@login_required
def api_check_bank():
    data = request.json or {}
    bank_id = data.get("bank_id")
    total_needed = int(data.get("total_questions", 10))
    if not bank_id:
        return jsonify({"success": False, "message": "Vui lòng chọn ngân hàng câu hỏi."}), 400

    bank = db.session.get(QuestionBank, bank_id)
    if not bank:
        return jsonify({"success": False, "message": "Ngân hàng không tồn tại."}), 404

    available_total = bank.question_count
    diff_stats = bank.stats_by_difficulty()
    type_stats = bank.stats_by_type()

    is_sufficient = available_total >= total_needed

    return jsonify({
        "success": True,
        "bank_name": bank.name,
        "is_sufficient": is_sufficient,
        "available_total": available_total,
        "total_needed": total_needed,
        "diff_stats": diff_stats,
        "type_stats": type_stats
    })


@app.route("/exams/<int:exam_id>/edit", methods=["GET", "POST"])
@login_required
def edit_exam(exam_id):
    exam = db.session.get(Exam, exam_id) or abort(404)
    if current_user.role == "teacher" and exam.created_by != current_user.id:
        abort(403)
    if current_user.role not in ("admin", "manager", "teacher"):
        abort(403)

    if current_user.role == "teacher":
        courses_list = Course.query.filter_by(teacher_id=current_user.id).order_by(Course.course_name).all()
    else:
        courses_list = Course.query.order_by(Course.course_name).all()

    if request.method == "POST":
        exam.title = request.form["title"]
        exam.exam_category = request.form.get("exam_category", "Kiểm tra thường kỳ")
        exam.course_id = int(request.form["course_id"])
        exam.duration_minutes = int(request.form.get("duration_minutes", 30))
        exam.shuffle_questions = bool(request.form.get("shuffle_questions"))
        exam.shuffle_answers = bool(request.form.get("shuffle_answers"))
        exam.is_active = bool(request.form.get("is_active"))
        start_time_str = request.form.get("start_time")
        end_time_str = request.form.get("end_time")
        exam.start_time = dt.datetime.strptime(start_time_str, "%Y-%m-%dT%H:%M") if start_time_str else None
        exam.end_time = dt.datetime.strptime(end_time_str, "%Y-%m-%dT%H:%M") if end_time_str else None

        # Delete old questions
        for q in exam.questions:
            db.session.delete(q)
        db.session.flush()

        # Re-add questions
        q_idx = 0
        while f"question_{q_idx}_text" in request.form:
            q_text = request.form[f"question_{q_idx}_text"]
            q_type = request.form.get(f"question_{q_idx}_type", "single")
            q_points = float(request.form.get(f"question_{q_idx}_points", 1))
            if q_text.strip():
                question = ExamQuestion(
                    exam_id=exam.id,
                    question_text=q_text,
                    question_type=q_type,
                    points=q_points,
                    order_index=q_idx,
                )
                db.session.add(question)
                db.session.flush()

                a_idx = 0
                while f"question_{q_idx}_answer_{a_idx}_text" in request.form:
                    a_text = request.form[f"question_{q_idx}_answer_{a_idx}_text"]
                    if a_text.strip():
                        if q_type == "truefalse":
                            is_correct = request.form.get(f"question_{q_idx}_correct") == str(a_idx)
                        elif q_type == "multi":
                            is_correct = request.form.get(f"question_{q_idx}_answer_{a_idx}_correct") == "1"
                        else:
                            is_correct = request.form.get(f"question_{q_idx}_correct") == str(a_idx)
                        answer = ExamAnswer(
                            question_id=question.id,
                            answer_text=a_text,
                            is_correct=is_correct,
                        )
                        db.session.add(answer)
                    a_idx += 1
            q_idx += 1

        db.session.commit()
        flash("Đã cập nhật bài kiểm tra.", "success")
        return redirect(url_for("exam_list"))

    return render_template("exam_form.html", exam=exam, courses=courses_list)


@app.route("/exams/<int:exam_id>/delete", methods=["POST"])
@login_required
def delete_exam(exam_id):
    exam = db.session.get(Exam, exam_id) or abort(404)
    if current_user.role == "teacher" and exam.created_by != current_user.id:
        abort(403)
    if current_user.role not in ("admin", "manager", "teacher"):
        abort(403)
    db.session.delete(exam)
    db.session.commit()
    flash("Đã xóa bài kiểm tra.", "success")
    return redirect(url_for("exam_list"))


@app.route("/exams/<int:exam_id>/take", methods=["GET", "POST"])
@login_required
def take_exam(exam_id):
    if current_user.role != "student":
        flash("Chỉ học sinh mới được làm bài kiểm tra.", "danger")
        return redirect(url_for("exam_list"))

    exam = db.session.get(Exam, exam_id) or abort(404)

    # Check student is enrolled
    course_ids = [e.course_id for e in Enrollment.query.filter_by(student_id=current_user.id, status="active").all()]
    if exam.course_id not in course_ids:
        flash("Bạn không có quyền làm bài kiểm tra này.", "danger")
        return redirect(url_for("exam_list"))

    # Check exam status
    if exam.status_label != "open":
        flash("Bài kiểm tra chưa mở hoặc đã đóng.", "warning")
        return redirect(url_for("exam_list"))

    # Check existing submission
    existing = ExamSubmission.query.filter_by(exam_id=exam.id, student_id=current_user.id).first()
    if existing and existing.is_graded:
        flash("Bạn đã làm bài kiểm tra này rồi.", "info")
        return redirect(url_for("exam_result", exam_id=exam.id))

    questions = list(exam.questions)
    if exam.shuffle_questions:
        random.shuffle(questions)

    if exam.shuffle_answers:
        for q in questions:
            q._shuffled_answers = list(q.answers)
            random.shuffle(q._shuffled_answers)
        else:
            for q in questions:
                q._shuffled_answers = list(q.answers)
    else:
        for q in questions:
            q._shuffled_answers = list(q.answers)

    if request.method == "POST":
        submission = existing or ExamSubmission(
            exam_id=exam.id,
            student_id=current_user.id,
            started_at=dt.datetime.utcnow(),
        )
        if not existing:
            db.session.add(submission)
            db.session.flush()

        total_score = 0
        max_score = 0

        for question in exam.questions:
            max_score += question.points
            q_type = question.question_type

            if q_type in ("truefalse", "true_false"):
                answers = question.answers
                total_statements = len(answers)
                correct_statements = 0
                selected_dict = {}

                for a in answers:
                    choice = request.form.get(f"q_{question.id}_ans_{a.id}")
                    if choice == "1":
                        selected_dict[str(a.id)] = True
                        if a.is_correct:
                            correct_statements += 1
                    elif choice == "0":
                        selected_dict[str(a.id)] = False
                        if not a.is_correct:
                            correct_statements += 1
                    else:
                        selected_dict[str(a.id)] = None

                if total_statements > 0:
                    if correct_statements == total_statements:
                        points_earned = question.points
                    elif correct_statements == 3:
                        points_earned = round(question.points * 0.5, 2)
                    elif correct_statements == 2:
                        points_earned = round(question.points * 0.25, 2)
                    elif correct_statements == 1:
                        points_earned = round(question.points * 0.1, 2)
                    else:
                        points_earned = 0
                else:
                    points_earned = 0

                is_correct = (correct_statements == total_statements) and (total_statements > 0)
                total_score += points_earned

                SubmissionAnswer.query.filter_by(submission_id=submission.id, question_id=question.id).delete()
                sa = SubmissionAnswer(
                    submission_id=submission.id,
                    question_id=question.id,
                    selected_answer_ids=json.dumps([a.id for a in answers if selected_dict.get(str(a.id)) is True]),
                    answer_text=json.dumps(selected_dict),
                    is_correct=is_correct,
                    points_earned=points_earned,
                )
                db.session.add(sa)

            elif q_type in ("short_answer", "essay", "tu_luan"):
                text_ans = request.form.get(f"q_{question.id}_text", "").strip()
                file_obj = request.files.get(f"q_{question.id}_file")
                file_url = None

                if file_obj and file_obj.filename:
                    upload_folder = os.path.join(app.static_folder, "uploads", "submissions")
                    os.makedirs(upload_folder, exist_ok=True)
                    ext = os.path.splitext(file_obj.filename)[1]
                    unique_filename = f"sub_{submission.id}_q_{question.id}_{int(time.time())}{ext}"
                    filepath = os.path.join(upload_folder, unique_filename)
                    file_obj.save(filepath)
                    file_url = url_for("static", filename=f"uploads/submissions/{unique_filename}")

                ans_payload = {"text": text_ans, "file_url": file_url}

                points_earned = 0
                is_correct = False
                correct_answers = [a.answer_text.strip().lower() for a in question.answers if a.answer_text]
                if text_ans and text_ans.lower() in correct_answers:
                    is_correct = True
                    points_earned = question.points
                    total_score += points_earned

                SubmissionAnswer.query.filter_by(submission_id=submission.id, question_id=question.id).delete()
                sa = SubmissionAnswer(
                    submission_id=submission.id,
                    question_id=question.id,
                    selected_answer_ids="[]",
                    answer_text=json.dumps(ans_payload),
                    is_correct=is_correct,
                    points_earned=points_earned,
                )
                db.session.add(sa)

            elif q_type == "multi":
                selected_ids = [int(x) for x in request.form.getlist(f"q_{question.id}")]
                correct_ids = {a.id for a in question.answers if a.is_correct}
                selected_set = set(selected_ids)
                is_correct = selected_set == correct_ids and len(selected_ids) > 0
                points_earned = question.points if is_correct else 0
                total_score += points_earned

                SubmissionAnswer.query.filter_by(submission_id=submission.id, question_id=question.id).delete()
                sa = SubmissionAnswer(
                    submission_id=submission.id,
                    question_id=question.id,
                    selected_answer_ids=json.dumps(selected_ids),
                    is_correct=is_correct,
                    points_earned=points_earned,
                )
                db.session.add(sa)

            else:
                val = request.form.get(f"q_{question.id}")
                selected_ids = [int(val)] if val else []
                correct_ids = {a.id for a in question.answers if a.is_correct}
                selected_set = set(selected_ids)
                is_correct = selected_set == correct_ids and len(selected_ids) > 0
                points_earned = question.points if is_correct else 0
                total_score += points_earned

                SubmissionAnswer.query.filter_by(submission_id=submission.id, question_id=question.id).delete()
                sa = SubmissionAnswer(
                    submission_id=submission.id,
                    question_id=question.id,
                    selected_answer_ids=json.dumps(selected_ids),
                    is_correct=is_correct,
                    points_earned=points_earned,
                )
                db.session.add(sa)

        has_manual_grading = any(q.question_type in ("short_answer", "essay", "tu_luan") for q in exam.questions)

        submission.total_score = total_score
        submission.max_score = max_score
        submission.is_graded = not has_manual_grading
        submission.submitted_at = dt.datetime.utcnow()
        db.session.commit()

        # Update student performance profile
        try:
            AdaptiveEngine.update_performance(current_user.id, submission)
        except Exception as e:
            print(f"Error updating performance profile: {e}")

        flash(f"Đã nộp bài! Điểm: {total_score}/{max_score}", "success")
        return redirect(url_for("exam_result", exam_id=exam.id))

    return render_template("exam_take.html", exam=exam, questions=questions)


@app.route("/exams/<int:exam_id>/result")
@login_required
def exam_result(exam_id):
    exam = db.session.get(Exam, exam_id) or abort(404)

    if current_user.role == "student":
        submission = ExamSubmission.query.filter_by(exam_id=exam.id, student_id=current_user.id).first()
        if not submission:
            flash("Bạn chưa làm bài kiểm tra này.", "warning")
            return redirect(url_for("exam_list"))
        return render_template("exam_result.html", exam=exam, submission=submission, single_view=True)

    if current_user.role in ("admin", "manager"):
        submissions = ExamSubmission.query.filter_by(exam_id=exam.id).order_by(ExamSubmission.total_score.desc()).all()
    elif current_user.role == "teacher":
        if exam.created_by != current_user.id:
            abort(403)
        submissions = ExamSubmission.query.filter_by(exam_id=exam.id).order_by(ExamSubmission.total_score.desc()).all()
    else:
        abort(403)

    return render_template("exam_result.html", exam=exam, submissions=submissions, single_view=False)


@app.route("/submissions/<int:submission_id>/grade", methods=["GET", "POST"])
@login_required
@teacher_manager_admin_required
def grade_submission(submission_id):
    submission = db.session.get(ExamSubmission, submission_id) or abort(404)
    exam = submission.exam
    
    if current_user.role == "teacher" and exam.created_by != current_user.id:
        abort(403)

    if request.method == "POST":
        total_score = 0.0
        
        for q in exam.questions:
            # Check if this question's grade was submitted
            grade_key = f"grade_q_{q.id}"
            if grade_key in request.form:
                try:
                    points = float(request.form[grade_key])
                    # Ensure points don't exceed max question points
                    points = max(0.0, min(points, q.points))
                except ValueError:
                    points = 0.0
                
                # Update SubmissionAnswer
                sa = SubmissionAnswer.query.filter_by(submission_id=submission.id, question_id=q.id).first()
                if not sa:
                    sa = SubmissionAnswer(
                        submission_id=submission.id,
                        question_id=q.id,
                        selected_answer_ids="[]",
                        is_correct=False,
                        points_earned=points
                    )
                    db.session.add(sa)
                else:
                    sa.points_earned = points
                
                total_score += points

        submission.total_score = total_score
        submission.is_graded = True
        db.session.commit()
        
        flash(f"Đã lưu điểm cho học sinh {submission.student.fullname}.", "success")
        return redirect(url_for("exam_result", exam_id=exam.id))

    return render_template("grade_submission.html", exam=exam, submission=submission)


# ──────────────────────── ADAPTIVE PRACTICE & ANALYTICS ────────────────────────

@app.route("/practice")
@login_required
def practice_home():
    if current_user.role != "student":
        flash("Tính năng tự luyện dành cho học sinh.", "info")
        return redirect(url_for("index"))

    # Find subjects for enrolled courses
    enrollments = Enrollment.query.filter_by(student_id=current_user.id, status="active").all()
    subjects = list({e.course.subject for e in enrollments if e.course and e.course.subject})
    if not subjects:
        subjects = ["Toán", "Tiếng Anh", "Vật Lý"]

    profiles = {}
    for sub in subjects:
        profiles[sub] = AdaptiveEngine.analyze_performance(current_user.id, sub)

    # Get available banks per subject
    banks_by_subject = {}
    for sub in subjects:
        banks_by_subject[sub] = QuestionBank.query.filter_by(subject=sub, is_active=True).all()

    return render_template("practice_home.html", subjects=subjects, profiles=profiles, banks_by_subject=banks_by_subject)


@app.route("/practice/generate", methods=["POST"])
@login_required
def practice_generate():
    if current_user.role != "student":
        abort(403)

    subject = request.form.get("subject", "Toán")
    bank_id = request.form.get("bank_id")
    total_q = int(request.form.get("total_questions", 10))

    if not bank_id:
        bank = QuestionBank.query.filter_by(subject=subject, is_active=True).first()
        if not bank:
            flash(f"Chưa có ngân hàng câu hỏi nào cho môn {subject}.", "warning")
            return redirect(url_for("practice_home"))
        bank_id = bank.id
    else:
        bank_id = int(bank_id)
        bank = db.session.get(QuestionBank, bank_id)

    # Get adaptive recommendation
    profile = AdaptiveEngine.analyze_performance(current_user.id, subject)
    recommendation = AdaptiveEngine.recommend_difficulty(profile)
    diff_cfg = recommendation["diff_config"]
    type_cfg = recommendation["type_config"]

    # Dummy course for practice exam reference if needed
    enrollment = Enrollment.query.filter_by(student_id=current_user.id, status="active").first()
    course_id = enrollment.course_id if enrollment else 1

    # Create Practice Exam
    exam = Exam(
        title=f"Bài Tự Luyện: {bank.name} ({recommendation['level'].upper()})",
        course_id=course_id,
        created_by=current_user.id,
        duration_minutes=int(total_q * 1.5),
        exam_type="practice",
        question_bank_id=bank.id,
        total_questions=total_q,
        difficulty_config=json.dumps(diff_cfg),
        question_type_config=json.dumps(type_cfg),
        is_active=True
    )
    db.session.add(exam)
    db.session.flush()

    selected_bqs = _select_random_bank_questions(bank.id, total_q, diff_cfg, type_cfg)

    for idx, bq in enumerate(selected_bqs):
        eq = ExamQuestion(
            exam_id=exam.id,
            bank_question_id=bq.id,
            question_text=bq.question_text,
            question_type=bq.question_type,
            difficulty_level=bq.difficulty_level,
            explanation=bq.explanation,
            points=bq.points,
            order_index=idx
        )
        db.session.add(eq)
        db.session.flush()

        for opt in bq.options:
            ea = ExamAnswer(
                question_id=eq.id,
                answer_text=opt.option_text,
                is_correct=opt.is_correct
            )
            db.session.add(ea)

    db.session.commit()
    return redirect(url_for("take_exam", exam_id=exam.id))


@app.route("/practice/progress")
@login_required
def practice_progress():
    if current_user.role != "student":
        flash("Trang này dành cho học sinh.", "info")
        return redirect(url_for("index"))

    enrollments = Enrollment.query.filter_by(student_id=current_user.id, status="active").all()
    subjects = list({e.course.subject for e in enrollments if e.course and e.course.subject})
    if not subjects:
        subjects = ["Toán", "Tiếng Anh"]

    selected_subject = request.args.get("subject", subjects[0] if subjects else "Toán")
    profile = AdaptiveEngine.analyze_performance(current_user.id, selected_subject)

    # Submission history
    subs = (
        ExamSubmission.query
        .filter_by(student_id=current_user.id)
        .order_by(ExamSubmission.submitted_at.asc())
        .all()
    )
    history = [
        {
            "date": s.submitted_at.strftime("%d/%m/%Y %H:%M") if s.submitted_at else "",
            "score": round(s.percentage / 10.0, 1),
            "title": s.exam.title if s.exam else "Bài luyện",
            "type": s.exam_type
        }
        for s in subs if s.exam and s.exam.course and s.exam.course.subject == selected_subject
    ]

    return render_template(
        "practice_progress.html",
        subjects=subjects,
        selected_subject=selected_subject,
        profile=profile,
        history=history
    )


@app.route("/analytics/teacher")
@login_required
@teacher_manager_admin_required
def teacher_analytics():
    if current_user.role == "teacher":
        my_courses = Course.query.filter_by(teacher_id=current_user.id).all()
        course_ids = [c.id for c in my_courses]
    else:
        my_courses = Course.query.all()
        course_ids = [c.id for c in my_courses]

    enrollments = Enrollment.query.filter(Enrollment.course_id.in_(course_ids), Enrollment.status == "active").all() if course_ids else []
    student_ids = list({e.student_id for e in enrollments})
    students = User.query.filter(User.id.in_(student_ids)).order_by(User.fullname).all() if student_ids else []

    selected_student_id = request.args.get("student_id", type=int)
    student_analytics = None
    if selected_student_id:
        st = db.session.get(User, selected_student_id)
        if st:
            profiles = {}
            for c in my_courses:
                profiles[c.subject] = AdaptiveEngine.analyze_performance(st.id, c.subject)
            student_analytics = {"student": st, "profiles": profiles}

    return render_template(
        "teacher_analytics.html",
        courses=my_courses,
        students=students,
        selected_student_id=selected_student_id,
        student_analytics=student_analytics
    )


@app.route("/analytics/parent")
@login_required
def parent_dashboard():
    if current_user.role != "parent" and not current_user.is_admin:
        abort(403)

    student = current_user.linked_student
    if not student:
        flash("Tài khoản phụ huynh chưa được liên kết với học sinh.", "warning")
        return redirect(url_for("index"))

    enrollments = Enrollment.query.filter_by(student_id=student.id, status="active").all()
    subjects = list({e.course.subject for e in enrollments if e.course and e.course.subject})
    if not subjects:
        subjects = ["Toán", "Tiếng Anh"]

    profiles = {}
    for sub in subjects:
        profiles[sub] = AdaptiveEngine.analyze_performance(student.id, sub)

    recent_submissions = (
        ExamSubmission.query
        .filter_by(student_id=student.id)
        .order_by(ExamSubmission.submitted_at.desc())
        .limit(10)
        .all()
    )

    return render_template(
        "parent_dashboard.html",
        student=student,
        subjects=subjects,
        profiles=profiles,
        recent_submissions=recent_submissions
    )


# ──────────────────────── SEED / MIGRATE ────────────────────────

def seed_sample_data():
    if User.query.count():
        return

    users = [
        ("admin", "Nguyễn Văn Admin", "ADMIN001", "admin", "0912345678", "Giám đốc", 0),
        ("manager", "Trần Thị Quản Lý", "QL001", "manager", "0987654321", "Quản lý học vụ", 0),
        ("gvtoan", "Cô Lan - Toán", "GV001", "teacher", "0938123456", "Giáo viên Toán", 150000),
        ("gvanh", "Thầy Minh - Anh", "GV002", "teacher", "0978123456", "Giáo viên Anh", 200000),
        ("hs001", "Nguyễn Thị Hoa", "HS001", "student", "0123456789", "Lớp 9", 0),
        ("hs002", "Trần Văn An", "HS002", "student", "0123456790", "Lớp 10", 0),
        ("ph001", "Phụ huynh Hoa", "PH001", "parent", "0987000000", "Phụ huynh", 0),
    ]
    created = {}
    for username, fullname, code, role, phone, position, rate in users:
        user = User(
            username=username,
            fullname=fullname,
            user_code=code,
            role=role,
            phone=phone,
            position=position,
            salary_rate_per_hour=rate,
            is_active=True,
        )
        user.set_password("123456")
        db.session.add(user)
        created[username] = user
    db.session.flush()
    created["ph001"].linked_student_id = created["hs001"].id

    courses_data = [
        ("Toán 9 nâng cao", "Toán", created["gvtoan"], "P101", 1200000, "monthly", "2026-06-01", "2026-12-31"),
        ("Anh văn 10 giao tiếp", "Anh văn", created["gvanh"], "P102", 5000000, "full_course", "2026-06-15", "2026-10-15"),
    ]
    courses_created = []
    for name, subject, teacher, room, tuition, tuition_type, start, end in courses_data:
        course = Course(
            course_name=name,
            subject=subject,
            teacher_id=teacher.id,
            classroom=room,
            description=f"Chương trình {name} tại trung tâm.",
            tuition_amount=tuition,
            tuition_type=tuition_type,
            start_date=parse_date(start),
            end_date=parse_date(end),
            is_active=True,
        )
        db.session.add(course)
        courses_created.append(course)
    db.session.flush()

    schedule_rows = [
        (courses_created[0], created["gvtoan"], "P101", 0, "18:00", "19:30"),
        (courses_created[0], created["gvtoan"], "P101", 2, "18:00", "20:00"),
        (courses_created[1], created["gvanh"], "P102", 5, "08:00", "10:00"),
    ]
    for course, teacher, room, weekday, start_value, end_value in schedule_rows:
        start = parse_time(start_value)
        end = parse_time(end_value)
        db.session.add(
            Schedule(
                course_id=course.id,
                teacher_id=teacher.id,
                classroom=room,
                weekday=weekday,
                start_time=start,
                end_time=end,
                duration_hours=duration_hours(start, end),
            )
        )

    db.session.add(Enrollment(student_id=created["hs001"].id, course_id=courses_created[0].id))
    db.session.add(Enrollment(student_id=created["hs002"].id, course_id=courses_created[1].id))
    db.session.add(
        TuitionPayment(
            student_id=created["hs001"].id,
            course_id=courses_created[0].id,
            amount=1200000,
            payment_type="monthly",
            payment_date=dt.date.today(),
            status="paid",
            note="Học phí tháng hiện tại",
        )
    )
    db.session.commit()


def table_columns(table_name):
    try:
        inspector = db.inspect(db.engine)
        columns = inspector.get_columns(table_name)
        return {col['name'] for col in columns}
    except Exception as e:
        print(f"Error inspecting table {table_name}: {e}")
        return set()


def add_column_if_missing(table_name, column_name, ddl):
    if column_name not in table_columns(table_name):
        # Quote table name to prevent conflicts with reserved keywords (like "user" in Postgres)
        db.session.execute(db.text(f'ALTER TABLE "{table_name}" ADD COLUMN {column_name} {ddl}'))
        db.session.commit()


def migrate_existing_sqlite_schema():
    inspector = db.inspect(db.engine)
    tables = set(inspector.get_table_names())

    # 1. SQLite specific table recreation to fix NOT NULL constraint mismatches
    if db.engine.dialect.name == "sqlite":
        # Rebuild course table if it contains the old 'name' column or 'image_file'
        if "course" in tables:
            columns = table_columns("course")
            if "name" in columns or "image_file" in columns:
                try:
                    db.session.execute(db.text("PRAGMA foreign_keys=OFF;"))
                    db.session.execute(db.text("""
                        CREATE TABLE course_new (
                            id INTEGER NOT NULL PRIMARY KEY AUTOINCREMENT,
                            course_name VARCHAR(120) NOT NULL,
                            subject VARCHAR(80) NOT NULL,
                            teacher_id INTEGER,
                            classroom VARCHAR(40),
                            description TEXT,
                            tuition_amount INTEGER DEFAULT 0,
                            tuition_type VARCHAR(20) DEFAULT 'monthly',
                            start_date DATE,
                            end_date DATE,
                            is_active BOOLEAN DEFAULT 1,
                            FOREIGN KEY(teacher_id) REFERENCES user (id)
                        );
                    """))
                    db.session.execute(db.text("""
                        INSERT INTO course_new (id, course_name, subject, teacher_id, classroom, description, tuition_amount, tuition_type, start_date, end_date, is_active)
                        SELECT 
                            id, 
                            COALESCE(course_name, name, ''), 
                            COALESCE(subject, ''), 
                            teacher_id, 
                            classroom, 
                            COALESCE(description, summary), 
                            COALESCE(tuition_amount, fee_per_session, 0), 
                            COALESCE(tuition_type, 'monthly'), 
                            start_date, 
                            end_date, 
                            COALESCE(is_active, 1)
                        FROM course;
                    """))
                    db.session.execute(db.text("DROP TABLE course;"))
                    db.session.execute(db.text("ALTER TABLE course_new RENAME TO course;"))
                    db.session.commit()
                except Exception as e:
                    db.session.rollback()
                    print(f"Error rebuilding SQLite course table: {e}")
                finally:
                    db.session.execute(db.text("PRAGMA foreign_keys=ON;"))

        # Rebuild schedule table if it contains the old 'day_of_week' column
        if "schedule" in tables:
            columns = table_columns("schedule")
            if "day_of_week" in columns:
                try:
                    db.session.execute(db.text("PRAGMA foreign_keys=OFF;"))
                    db.session.execute(db.text("""
                        CREATE TABLE schedule_new (
                            id INTEGER NOT NULL PRIMARY KEY AUTOINCREMENT,
                            course_id INTEGER NOT NULL,
                            teacher_id INTEGER NOT NULL,
                            classroom VARCHAR(40) NOT NULL,
                            weekday INTEGER NOT NULL,
                            start_time TIME NOT NULL,
                            end_time TIME NOT NULL,
                            duration_hours FLOAT NOT NULL,
                            FOREIGN KEY(course_id) REFERENCES course (id) ON DELETE CASCADE,
                            FOREIGN KEY(teacher_id) REFERENCES user (id)
                        );
                    """))
                    db.session.execute(db.text("""
                        INSERT INTO schedule_new (id, course_id, teacher_id, classroom, weekday, start_time, end_time, duration_hours)
                        SELECT 
                            id, 
                            course_id, 
                            teacher_id, 
                            COALESCE(classroom, room, ''), 
                            COALESCE(weekday, day_of_week - 1, 0), 
                            start_time, 
                            end_time, 
                            COALESCE(duration_hours, 1.5)
                        FROM schedule;
                    """))
                    db.session.execute(db.text("DROP TABLE schedule;"))
                    db.session.execute(db.text("ALTER TABLE schedule_new RENAME TO schedule;"))
                    db.session.commit()
                except Exception as e:
                    db.session.rollback()
                    print(f"Error rebuilding SQLite schedule table: {e}")
                finally:
                    db.session.execute(db.text("PRAGMA foreign_keys=ON;"))

        # Rebuild attendance table if it contains the old 'teacher_id' column
        if "attendance" in tables:
            columns = table_columns("attendance")
            if "teacher_id" in columns:
                try:
                    db.session.execute(db.text("PRAGMA foreign_keys=OFF;"))
                    db.session.execute(db.text("""
                        CREATE TABLE attendance_new (
                            id INTEGER NOT NULL PRIMARY KEY AUTOINCREMENT,
                            schedule_id INTEGER,
                            student_id INTEGER NOT NULL,
                            course_id INTEGER NOT NULL,
                            date DATE NOT NULL,
                            status VARCHAR(20) NOT NULL DEFAULT 'present',
                            note TEXT,
                            marked_by_id INTEGER,
                            created_at DATETIME,
                            updated_at DATETIME,
                            CONSTRAINT uq_attendance_student_course_date UNIQUE (student_id, course_id, date),
                            FOREIGN KEY(student_id) REFERENCES user (id),
                            FOREIGN KEY(course_id) REFERENCES course (id),
                            FOREIGN KEY(marked_by_id) REFERENCES user (id)
                        );
                    """))
                    db.session.execute(db.text("""
                        INSERT INTO attendance_new (id, schedule_id, student_id, course_id, date, status, note, marked_by_id, created_at, updated_at)
                        SELECT 
                            id, 
                            schedule_id, 
                            COALESCE(student_id, 1), 
                            COALESCE(course_id, 1), 
                            date, 
                            COALESCE(status, 'present'), 
                            note, 
                            COALESCE(marked_by_id, teacher_id), 
                            created_at, 
                            updated_at
                        FROM attendance;
                    """))
                    db.session.execute(db.text("DROP TABLE attendance;"))
                    db.session.execute(db.text("ALTER TABLE attendance_new RENAME TO attendance;"))
                    db.session.commit()
                except Exception as e:
                    db.session.rollback()
                    print(f"Error rebuilding SQLite attendance table: {e}")
                finally:
                    db.session.execute(db.text("PRAGMA foreign_keys=ON;"))

    # 2. General check & alter DDL for added/missing columns (DB agnostic)
    if "user" in tables:
        add_column_if_missing("user", "salary_rate_per_hour", "INTEGER DEFAULT 0")
        add_column_if_missing("user", "linked_student_id", "INTEGER")

    if "course" in tables:
        add_column_if_missing("course", "course_name", "VARCHAR(120) DEFAULT ''")
        add_column_if_missing("course", "classroom", "VARCHAR(40)")
        add_column_if_missing("course", "description", "TEXT")
        add_column_if_missing("course", "tuition_amount", "INTEGER DEFAULT 0")
        add_column_if_missing("course", "tuition_type", "VARCHAR(20) DEFAULT 'monthly'")
        add_column_if_missing("course", "start_date", "DATE")
        add_column_if_missing("course", "end_date", "DATE")
        add_column_if_missing("course", "is_active", "BOOLEAN DEFAULT TRUE")
        columns = table_columns("course")
        if "name" in columns:
            db.session.execute(db.text("UPDATE course SET course_name = name WHERE course_name IS NULL OR course_name = ''"))
        if "summary" in columns:
            db.session.execute(db.text("UPDATE course SET description = summary WHERE description IS NULL"))
        if "fee_per_session" in columns:
            db.session.execute(db.text("UPDATE course SET tuition_amount = fee_per_session WHERE tuition_amount IS NULL OR tuition_amount = 0"))
        db.session.commit()

    if "schedule" in tables:
        add_column_if_missing("schedule", "weekday", "INTEGER DEFAULT 0")
        add_column_if_missing("schedule", "classroom", "VARCHAR(40) DEFAULT ''")
        add_column_if_missing("schedule", "duration_hours", "FLOAT DEFAULT 1.5")
        columns = table_columns("schedule")
        if "day_of_week" in columns:
            db.session.execute(db.text("UPDATE schedule SET weekday = CASE WHEN day_of_week BETWEEN 1 AND 7 THEN day_of_week - 1 ELSE 0 END"))
        if "room" in columns:
            db.session.execute(db.text("UPDATE schedule SET classroom = room WHERE classroom IS NULL OR classroom = ''"))
        db.session.commit()

    if "attendance" in tables:
        add_column_if_missing("attendance", "course_id", "INTEGER DEFAULT 1")
        add_column_if_missing("attendance", "marked_by_id", "INTEGER")
        add_column_if_missing("attendance", "updated_at", "DATETIME")
        columns = table_columns("attendance")
        if "teacher_id" in columns:
            db.session.execute(db.text("UPDATE attendance SET marked_by_id = teacher_id WHERE marked_by_id IS NULL"))
        if "schedule_id" in columns and "schedule" in tables:
            db.session.execute(
                db.text(
                    "UPDATE attendance SET course_id = "
                    "(SELECT schedule.course_id FROM schedule WHERE schedule.id = attendance.schedule_id) "
                    "WHERE course_id IS NULL OR course_id = 1"
                )
            )
        db.session.commit()

    if "teaching_record" in tables:
        add_column_if_missing("teaching_record", "confirmed_by_teacher", "BOOLEAN DEFAULT FALSE")
        add_column_if_missing("teaching_record", "confirmed_at", "DATETIME")

    if "bank_question" in tables:
        add_column_if_missing("bank_question", "context", "TEXT")
        add_column_if_missing("bank_question", "image_url", "VARCHAR(500)")
        add_column_if_missing("bank_question", "confidence_scores", "TEXT")

    # Migrate Exam tables — add new columns for question bank integration
    if "exam" in tables:
        add_column_if_missing("exam", "question_bank_id", "INTEGER")
        add_column_if_missing("exam", "exam_type", "VARCHAR(20) DEFAULT 'exam'")
        add_column_if_missing("exam", "exam_category", "VARCHAR(50) DEFAULT 'Kiểm tra thường kỳ'")
        add_column_if_missing("exam", "total_questions", "INTEGER")
        add_column_if_missing("exam", "difficulty_config", "TEXT")
        add_column_if_missing("exam", "question_type_config", "TEXT")

    if "exam_question" in tables:
        add_column_if_missing("exam_question", "bank_question_id", "INTEGER")
        add_column_if_missing("exam_question", "difficulty_level", "VARCHAR(20)")
        add_column_if_missing("exam_question", "explanation", "TEXT")
        add_column_if_missing("exam_question", "context", "TEXT")
        add_column_if_missing("exam_question", "image_url", "VARCHAR(500)")

    if "exam_submission" in tables:
        add_column_if_missing("exam_submission", "exam_type", "VARCHAR(20) DEFAULT 'exam'")
        add_column_if_missing("exam_submission", "time_spent_seconds", "INTEGER")
        add_column_if_missing("exam_submission", "correct_count", "INTEGER DEFAULT 0")
        add_column_if_missing("exam_submission", "wrong_count", "INTEGER DEFAULT 0")

    if "submission_answer" in tables:
        add_column_if_missing("submission_answer", "answer_text", "TEXT")
        add_column_if_missing("submission_answer", "time_spent_seconds", "INTEGER")

    if "enrollment" in tables:
        # Clean up string student_ids stored in the database to use integer IDs
        rows = db.session.execute(db.text("SELECT id, student_id FROM enrollment")).all()
        for r_id, s_id in rows:
            if isinstance(s_id, str) and not str(s_id).isdigit():
                # Look up the user by username or user_code
                user = User.query.filter(
                    (User.role == "student") &
                    ((User.username == s_id.lower()) | (User.user_code == s_id.upper()) | (User.username == s_id) | (User.user_code == s_id))
                ).first()
                if user:
                    db.session.execute(db.text("UPDATE enrollment SET student_id = :uid WHERE id = :eid"), {"uid": user.id, "eid": r_id})
        db.session.commit()

def backfill_teaching_records_from_attendance():
    pairs = (
        db.session.query(Attendance.course_id, Attendance.date)
        .filter(Attendance.status == "present")
        .distinct()
        .all()
    )
    for course_id, class_date in pairs:
        course = db.session.get(Course, course_id)
        if course:
            sync_teaching_record(course, class_date)
    db.session.commit()

def initialize_database():
    db.create_all()
    migrate_existing_sqlite_schema()
    backfill_teaching_records_from_attendance()
    seed_sample_data()

# Always run database initialization on startup (both local and Vercel environments)
with app.app_context():
    initialize_database()

if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5002)
